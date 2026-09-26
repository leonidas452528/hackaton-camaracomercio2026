"""Genera los 6 entregables del evento (Word) de Territorio Preparado.

Uso: python3 crear_entregables.py   (desde cualquier carpeta)
Cada cifra viene de docs/PROYECTO.md o de los Word de los artefactos 1 a 5.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
SALIDA = AQUI.parent
RAIZ = SALIDA.parents[1]
CAPTURAS = RAIZ / "entregables/presentacion-ampliada/assets"
RENDERS = RAIZ / "maqueta3d/deliverables/renders"

TINTA = RGBColor(0x1F, 0x29, 0x33)
TITULO = RGBColor(0x00, 0x3D, 0x4A)
ACENTO = "00614F"
SUAVE = "E6F2EF"

EQUIPO = "William Ortiz · Herlin Echeverry · Pablo Arango · Bryan Martínez Villamarín"
URL_DEMO = "https://leonidas452528.github.io/hackaton-camaracomercio2026/"
URL_REPO = "https://github.com/leonidas452528/hackaton-camaracomercio2026"

LINEAS_ROJAS = [
    "La herramienta propone y la autoridad decide: la activación de un espacio la decide la Secretaría de Gestión del Riesgo o el Consejo Municipal de Gestión del Riesgo (Ley 1523 de 2012, Decreto 2157 de 2017).",
    "No reemplaza evaluaciones estructurales (Ley 400 de 1997, NSR-10): solo registra si existe una evaluación vigente.",
    "No invade las competencias de Bomberos (Ley 1575 de 2012) ni emite conceptos de seguridad humana o contra incendios.",
    "Datos mínimos (Ley 1581 de 2012): solo conteos agregados, sin nombres ni documentos y sin datos de menores (Ley 1098 de 2006).",
    "Sin reconocimiento facial ni identificación individual.",
    "Criterios explicables: cada propuesta muestra las variables y reglas que usó.",
]

FUENTES = {
    "defensoria": "https://www.defensoria.gov.co/web/guest/-/defensoria-pide-atencion-urgente-para-336-familias-en-albergues-de-cali",
    "elpais": "https://www.elpais.com.co/cali/alcaldia-de-cali-entrega-balance-de-gestion-tras-el-sismo-mas-de-3400-edificaciones-evaluadas-y-45000-familias-damnificadas-2338.html",
    "semana": "https://www.semana.com/nacion/cali/articulo/tras-el-terremoto-cali-informa-cuales-escenarios-deportivos-vuelven-a-estar-al-servicio-de-los-atletas-desde-este-31-de-agosto/202628/",
    "cauca": "https://cwmas.com.co/cali/2026/02/24/rio-cauca-en-alerta-naranja-aumenta-el-caudal-en-el-oriente-de-cali/",
    "ocha": "https://www.infobae.com/colombia/2026/09/25/ocha-alerta-vacios-de-informacion-y-atencion-tras-terremoto-que-deja-486917-afectados-en-colombia/",
    "epou": "https://datos.cali.gov.co/dataset/epou-espacio-publico-efectivo",
    "nino": "https://ciudadregion.com/regiones/valle-del-cauca/cali/pico-fenomeno-nino-noviembre-valle-cauca-niveles-embalses",
    "esfera": "https://spherestandards.org/wp-content/uploads/Sphere-Handbook-2018-EN.pdf",
    "comunas": "https://www.datos.gov.co/dataset/Comunas-de-Santiago-de-Cali/dx2g-2mhm",
    "osm": "https://www.openstreetmap.org/copyright",
    "ccby": "https://creativecommons.org/licenses/by-sa/4.0/",
}


# ---------- utilidades de formato ----------
def sombrear(celda, color):
    tc = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc.append(shd)


def nuevo_documento():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
    base = doc.styles["Normal"]
    base.font.name = "Arial"
    base.font.size = Pt(10.5)
    base.font.color.rgb = TINTA
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    base.paragraph_format.space_after = Pt(6)
    base.paragraph_format.line_spacing = 1.15
    for nivel, tam in ((1, 16), (2, 13), (3, 11.5)):
        st = doc.styles[f"Heading {nivel}"]
        st.font.name = "Arial"
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.color.rgb = TITULO
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.paragraph_format.space_before = Pt(14 if nivel == 1 else 10)
        st.paragraph_format.space_after = Pt(6)
    return doc


def pie(doc, texto):
    for s in doc.sections:
        p = s.footer.paragraphs[0]
        p.text = texto
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x55, 0x5F, 0x66)


def portada(doc, numero, titulo, origen, criterio):
    p = doc.add_paragraph()
    r = p.add_run("Hackathon Smart City Expo Cali 2026 · RETO-01 Cali Activa")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x00, 0x61, 0x4F)
    r.bold = True
    p = doc.add_paragraph()
    r = p.add_run(f"Entregable {numero} · {titulo}")
    r.font.size = Pt(22)
    r.bold = True
    r.font.color.rgb = TITULO
    p = doc.add_paragraph()
    r = p.add_run("Territorio Preparado: espacios y ayudas coordinados ante emergencias")
    r.font.size = Pt(12.5)
    r.italic = True
    tabla(
        doc,
        None,
        [
            ["Equipo", EQUIPO],
            ["Reto", "RETO-01 Cali Activa: espacios públicos que se transforman para cuidar"],
            ["Owner del reto", "Secretaría de Gestión del Riesgo de Emergencias y Desastres · líder técnico: DATIC"],
            ["Fase del proceso", origen],
            ["Criterio de evaluación que evidencia", criterio],
            ["Fecha", "Sábado 26 de septiembre de 2026"],
            ["Prototipo público", URL_DEMO],
        ],
        anchos=(4.6, 12.0),
        encabezado=False,
    )


def h(doc, texto, nivel=1):
    doc.add_heading(texto, level=nivel)


def par(doc, texto, negrita_inicial=None):
    p = doc.add_paragraph()
    if negrita_inicial:
        p.add_run(negrita_inicial).bold = True
    p.add_run(texto)
    return p


def vinetas(doc, items, estilo="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=estilo)
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def tabla(doc, cabecera, filas, anchos=None, encabezado=True):
    ncol = len(cabecera) if cabecera else len(filas[0])
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if cabecera:
        fila = t.add_row().cells
        for i, c in enumerate(cabecera):
            fila[i].text = ""
            r = fila[i].paragraphs[0].add_run(c)
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.size = Pt(9.5)
            sombrear(fila[i], ACENTO)
    for f in filas:
        celdas = t.add_row().cells
        for i, c in enumerate(f):
            celdas[i].text = ""
            r = celdas[i].paragraphs[0].add_run(str(c))
            r.font.size = Pt(9.5)
            if not encabezado and i == 0:
                r.bold = True
                sombrear(celdas[i], SUAVE)
    if anchos:
        t.autofit = False
        for i, w in enumerate(anchos):
            t.columns[i].width = Cm(w)
        for fila in t.rows:
            for i, w in enumerate(anchos):
                fila.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def imagen(doc, ruta, ancho_cm, pie_texto):
    if not ruta.exists():
        raise FileNotFoundError(ruta)
    doc.add_picture(str(ruta), width=Cm(ancho_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(pie_texto)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.italic = True
        r.font.size = Pt(8.5)


def lineas_rojas(doc):
    h(doc, "Líneas rojas que respeta este entregable")
    vinetas(doc, LINEAS_ROJAS)


def fuentes(doc, claves, extra=()):
    h(doc, "Fuentes")
    nombres = {
        "defensoria": "Defensoría del Pueblo: 336 familias en albergues de Cali",
        "elpais": "El País: balance de la Alcaldía tras el sismo (RUD al 17 de septiembre)",
        "semana": "Semana: escenarios deportivos al 31 de agosto",
        "cauca": "CWMAS: río Cauca en alerta naranja (24 de febrero de 2026)",
        "ocha": "Infobae/OCHA: vacíos de información (25 de septiembre de 2026)",
        "epou": "DAPM/IDESC: espacio público efectivo (datos abiertos de Cali)",
        "nino": "Ciudad Región: pico de El Niño en noviembre de 2026 (IDEAM)",
        "esfera": "Manual Esfera 2018",
        "comunas": "Datos abiertos: comunas de Santiago de Cali",
        "osm": "OpenStreetMap: derechos de autor y atribución",
        "ccby": "Licencia CC BY-SA 4.0 de los datos derivados",
    }
    vinetas(doc, [(nombres[k] + ": ", FUENTES[k]) for k in claves] + list(extra))
    par(doc, "Bitácora completa con todas las fuentes: " + URL_REPO + "/blob/main/docs/PROYECTO.md")


def guardar(doc, nombre):
    ruta = SALIDA / nombre
    doc.save(ruta)
    print("OK", ruta.relative_to(RAIZ))


# ---------- 1. Alineación con el reto ----------
def alineacion():
    doc = nuevo_documento()
    portada(doc, 1, "Alineación con el reto", "Apertura y apropiación del reto · Artefactos 1 y 2 (23 y 25 de septiembre)", "Pertinencia frente al reto y comprensión del problema (20 %)")

    h(doc, "1. El reto oficial")
    par(doc, "¿Cómo podríamos transformar los espacios públicos de Cali en una red adaptable que, según el tipo de emergencia y las necesidades de la población, permita activar y operar de manera segura y coordinada funciones temporales de atención y recuperación, y facilite posteriormente el retorno del espacio a su uso cotidiano?", "Pregunta retadora: ")
    par(doc, "diseñar y validar un prototipo de baja fidelidad para caracterizar espacios, analizar su aptitud según el tipo de emergencia, identificar las adecuaciones necesarias, activar entidades y servicios, hacer seguimiento y apoyar el retorno a su uso cotidiano.", "Objetivo: ")
    par(doc, "UAE Gestión de Bienes y Servicios, Secretarías de Infraestructura, Deporte, Seguridad y Justicia, Salud Pública, Bienestar Social y Educación, Planeación, EMCALI y UAESP.", "Entidades a articular: ")

    h(doc, "2. La frase del equipo (artefacto 1)")
    par(doc, "«Queremos ayudar a las familias que deben refugiarse en espacios públicos de Cali durante una emergencia, y a la Secretaría de Gestión del Riesgo que los habilita, que enfrenta espacios que se activan sobre la marcha, sin saber si sirven para esa amenaza, qué les falta ni qué entidad responde, para que cada espacio se prepare antes, se active con criterios claros y vuelva a su uso cotidiano.»")

    h(doc, "3. Quiénes viven el reto")
    tabla(doc, ["Rol", "Actores"], [
        ["Quién vive el problema", "Familias damnificadas (con niños, adultos mayores, personas con discapacidad y migrantes); comunidades de Chiminangos I y II y Calimio Norte que montaron albergues propios; vecinos y usuarios habituales de parques y escenarios."],
        ["Quién decide", "Secretaría de Gestión del Riesgo de Emergencias y Desastres y Consejo Municipal de Gestión del Riesgo (Ley 1523 de 2012). Líder técnico: DATIC."],
        ["Quién puede ayudar", "Secretaría del Deporte, DAPM, EMCALI, UAESP, ICBF, Secretaría de Salud, Bienestar Social, Bomberos, Cruz Roja, Defensa Civil, juntas de acción comunal, Defensoría del Pueblo y CVC."],
    ], anchos=(4.2, 12.4))

    h(doc, "4. Lo que sabemos: evidencia con semáforo")
    tabla(doc, ["Evidencia", "Estado", "Fuente"], [
        ["Tras el sismo del 10 de agosto de 2026 se reportaron más de 30 familias y 143 personas alojadas en la Cancha de Hockey.", "Verde", "Portafolio de Retos, Alcaldía de Cali, v2.0"],
        ["La Defensoría encontró 336 familias (1.046 personas) en 3 albergues autogestionados en canchas de Chiminangos I, Chiminangos II y Calimio Norte.", "Verde", "Defensoría del Pueblo"],
        ["45.138 familias inscritas en el Registro Único de Damnificados (corte al 17 de septiembre) y 3.403 edificaciones evaluadas.", "Verde", "El País"],
        ["Al 31 de agosto, de 43 escenarios deportivos revisados, 12 estaban habilitados y 22 (51,2 %) esperaban inspección técnica.", "Verde", "Semana"],
        ["El 24 de febrero de 2026 el río Cauca llegó a 9,25 m y se declaró alerta naranja en el oriente.", "Verde", "CWMAS"],
        ["OCHA reporta cifras que no coinciden, familias sin información sobre inspecciones y fallas de agua y saneamiento en albergues.", "Verde", "Infobae/OCHA"],
        ["Cali tiene 1.970 espacios públicos efectivos inventariados con geometría.", "Verde", "DAPM/IDESC"],
        ["El IDEAM prevé un El Niño muy fuerte con pico en noviembre de 2026.", "Verde", "Ciudad Región"],
        ["La cifra de fallecidos en Cali difiere entre fuentes (105 frente a 154).", "Amarillo: no se usa", "Repositorio oficial de la Alcaldía"],
        ["Hoy no existe una ficha de aptitud de cada espacio según la amenaza.", "Rojo: supuesto", "Por validar con la Secretaría"],
        ["No está claro qué entidad responde por cada servicio de un albergue.", "Rojo: supuesto", "Por validar con las entidades"],
    ], anchos=(10.0, 2.6, 4.0))

    h(doc, "5. Hallazgos propios (lo que la ficha no dice)")
    vinetas(doc, [
        ("Albergues fuera del circuito oficial. ", "La comunidad montó 3 albergues autogestionados (336 familias, 1.046 personas) en Chiminangos I y II y Calimio Norte (Defensoría del Pueblo)."),
        ("Se refugiaron de un sismo en una zona inundable. ", "Esas canchas tienen amenaza alta de inundación, suelo licuable y quedan a entre 72 y 152 m del dique del río Cauca (cruce propio de las amenazas del POT, Acuerdo 0373 de 2014, con el espacio público del DAPM/IDESC)."),
        ("La aptitud depende de la amenaza. ", "En el oriente, sobre el Jarillón (comunas 6, 7, 13, 14 y 21), entre el 55 y el 92 % de los espacios públicos tiene amenaza alta de inundación (cruce propio de los 1.970 espacios con el POT)."),
        ("El retorno es un cuello de botella. ", "Al 31 de agosto, 22 de 43 escenarios (51,2 %) seguían esperando inspección técnica para volver a su uso (Semana)."),
        ("Falta caracterización y trazabilidad. ", "La Defensoría observó traslados de ayudas con destino desconocido y OCHA reporta cifras que no coinciden entre fuentes."),
    ])

    h(doc, "6. El reto acotado (artefacto 2)")
    tabla(doc, None, [
        ["Usuario institucional (owner)", "Secretaría de Gestión del Riesgo de Emergencias y Desastres de Cali. Líder técnico: DATIC."],
        ["Beneficiario ciudadano", "Familias que deben alojarse temporalmente en espacios públicos durante una emergencia, en especial niños, adultos mayores y personas con discapacidad (siempre en conteos agregados), y las comunidades que hoy improvisan albergues propios."],
        ["Territorio concreto", "Oriente de Cali sobre el Jarillón del río Cauca (comunas 6, 7, 13, 14 y 21), con foco en Chiminangos I, Chiminangos II y Calimio Norte, y los escenarios usados tras el sismo (Cancha de Hockey Miguel Calero, Diamante de Béisbol y Unidad Deportiva Jaime Aparicio)."],
        ["Momento crítico de fricción", "Las primeras horas después de la alerta, cuando hay que decidir qué espacio abrir sin saber si sirve para esa amenaza, qué le falta y quién responde; y el cierre, cuando el espacio debe volver a su uso sin una revisión ni un acta clara."],
        ["Falla del proceso actual", "Los espacios se habilitan sobre la marcha, sin una ficha de aptitud por amenaza. Las necesidades y los responsables se coordinan por teléfono, sin registro de brechas ni seguimiento. El retorno depende de inspecciones represadas. Mientras tanto, la comunidad se refugia por su cuenta, incluso en zonas inundables."],
        ["Hipótesis de solución No-Code", "Una aplicación web con datos abiertos que, dada una amenaza y un conteo agregado de población, descarte los espacios expuestos, proponga hasta 3 candidatos con criterios visibles, calcule necesidades según Esfera y el responsable de cada servicio, y exporte un borrador para que la autoridad decida."],
    ], anchos=(4.6, 12.0), encabezado=False)
    p = par(doc, "¿Cómo podríamos ayudar a la Secretaría de Gestión del Riesgo a preparar y activar los espacios públicos del oriente de Cali según el tipo de amenaza, con criterios visibles, necesidades calculadas y responsables claros, para que las familias no terminen refugiándose en zonas expuestas y cada espacio vuelva a su uso cotidiano?", "¿Cómo podríamos…? ")

    h(doc, "7. Cómo responde la propuesta a cada indicador de la ficha")
    tabla(doc, ["Indicador del reto", "Cómo lo atiende Territorio Preparado"], [
        ["Diferenciar la respuesta según el tipo de emergencia", "Reglas por amenaza: en inundación descarta cruces con las capas del POT; en sismo excluye licuación y corrimiento y exige inspección; sequía e incendios no preseleccionan sin evidencia."],
        ["Número de variables de preparación", "Huella cartográfica, cruces de amenaza, baños, agua, área cubierta, reservas de circulación y atención, y evidencia pendiente por espacio."],
        ["Capacidad de identificar brechas", "Necesidad calculada frente a existencia; si la existencia se desconoce, la brecha queda como «no calculable», no como cero."],
        ["Trazabilidad necesidad → servicio → entidad", "Cada servicio tiene un responsable propuesto y un estado local («Por medir», «En revisión»)."],
        ["Tiempo para generar una propuesta de activación", "El borrador se genera en la demo en minutos. Falta medirlo con usuarios reales."],
        ["Ciclo completo, de la preparación a la recuperación", "Preparación y activación implementadas; recuperación representada en la maqueta (uso cotidiano → emergencia → recuperación), con el acta de retorno como siguiente paso."],
    ], anchos=(5.2, 11.4))

    lineas_rojas(doc)
    fuentes(doc, ["defensoria", "elpais", "semana", "cauca", "ocha", "epou", "nino"])
    pie(doc, "Territorio Preparado · Entregable 1 · Alineación con el reto")
    guardar(doc, "01_Alineacion_con_el_reto.docx")


# ---------- 2. Arquitectura de la solución ----------
def arquitectura():
    doc = nuevo_documento()
    portada(doc, 2, "Arquitectura de la solución", "Sprint de solución · Artefacto 3: ideación y filtro (25 de septiembre)", "Uso pertinente de tecnología y datos (10 %) · Calidad y coherencia del prototipo (10 %) · Innovación (15 %)")

    h(doc, "1. De 8 ideas a una solución")
    tabla(doc, ["#", "Idea", "Flujo", "Herramienta"], [
        ["1", "Territorio Preparado: propone candidatos según la amenaza, calcula cuánto cabe y qué falta, y a quién le corresponde", "Amenaza + sector + conteo agregado → cruce POT, IDESC y Esfera → candidatos, necesidades, responsables y borrador", "Aplicación web (React + Leaflet)"],
        ["2", "Cali Lista: caracterización de espacios y de su entorno con puntaje de aptitud", "Visita de campo → ficha nivel A (espacio) + nivel B (entorno) → puntaje 60/40", "AppSheet"],
        ["3", "Tablero de brechas por entidad", "Necesidades Esfera − existencias → brecha → responsable → estado", "Looker Studio"],
        ["4", "Kit modular reutilizable + IoT sin cámaras", "Brecha de área → kits en bodega → instalación → sensores (SIMULADO)", "Figma"],
        ["5", "Aviso comunitario de puntos confirmados", "Apertura confirmada → aviso → difusión (SIMULADA)", "Glide"],
        ["6", "Checklist y acta de retorno", "Fin de la operación → checklist → acta → espacio disponible", "Softr"],
        ["7", "Adecuaciones permanentes priorizadas", "Historial de amenazas + uso probable → ranking → adecuaciones", "Looker Studio"],
        ["8", "Registro de albergues autogestionados", "Reporte agregado de la JAC → mapa → alerta si hay amenaza", "Glide"],
    ], anchos=(0.8, 5.6, 6.6, 3.6))
    tabla(doc, ["Cuadrante", "Ideas"], [
        ["Alto impacto + alta factibilidad", "Territorio Preparado (integra el tablero de brechas y el checklist de retorno)"],
        ["Alto impacto + baja factibilidad", "Kit modular + IoT; adecuaciones permanentes; registro de albergues autogestionados (exige un protocolo de datos con las JAC)"],
        ["Bajo impacto + alta factibilidad", "Aviso comunitario de puntos confirmados; Cali Lista como formulario aislado"],
    ], anchos=(5.0, 11.6))
    par(doc, "es la única idea que conecta todo el ciclo que pide el reto (caracterizar, decidir, activar, seguir y retornar) con datos reales de Cali y sin tocar datos personales. Las demás quedan como módulos: la caracterización de Cali Lista alimenta la ficha, el tablero muestra las brechas y el kit cubre el área faltante.", "Por qué la elegimos: ")

    h(doc, "2. Flujo funcional: entrada → proceso → salida")
    tabla(doc, ["Etapa", "Qué ocurre"], [
        ["Entrada", "Amenaza (inundación, sismo, sequía, incendio forestal o en edificación), sector (comuna o barrio) y conteo agregado de personas del escenario (SIMULADO)."],
        ["Proceso", "Cruce con 1.970 espacios DAPM/IDESC, 1.021 registros deportivos, amenazas del POT (inundación, licuación, remoción) y 1.090 puntos de daño del sismo (Copernicus EMSR916 e ICube-SERTIT). Descarte de expuestos, cálculo de área, agua, baños y área cubierta (Esfera) y responsable por servicio."],
        ["Salida", "Hasta 3 candidatos explicados, necesidades y brechas por entidad, borrador exportable (JSON) y aviso comunitario solo después de la confirmación de la autoridad."],
    ], anchos=(2.6, 14.0))

    h(doc, "3. Arquitectura por capas")
    tabla(doc, ["Capa", "Componentes", "Dónde está"], [
        ["1. Datos abiertos", "Espacio público efectivo y escenarios deportivos (DAPM/IDESC, CC BY-SA 4.0), amenazas del POT (Acuerdo 0373 de 2014), comunas y barrios, daños satelitales del sismo (Copernicus, SERTIT). Cada fuente tiene URL, licencia y SHA-256.", "prototipo/datos/raw, FUENTES.md, manifest.json"],
        ["2. Preparación de datos", "Scripts en Python y Node que cruzan geometrías, asignan comuna y barrio y quitan campos personales (visitadores, contactos, identificadores prediales). Comparación con los servicios WFS: 9 capas coincidieron el 25 de septiembre.", "prototipo/scripts, maqueta3d/scripts (npm run data)"],
        ["3. Reglas explicables", "Elegibilidad por amenaza, preselección de hasta 3 candidatos por distancia tras los descartes, necesidades Esfera y responsable por servicio. Cada regla se muestra en pantalla.", "maqueta3d/src/planning.ts"],
        ["4. Dimensionamiento", "Área cartográfica del polígono, menos reservas (no utilizable, circulación, atención, otros usos) → módulos que caben. Bloquea propuestas que exceden el balance.", "maqueta3d/src/calculateAcopioCapacity.ts"],
        ["5. Interfaz", "Mapa Leaflet con calles OpenStreetMap, fichas, enlaces a Google Maps y Street View, preparación de intervención, análisis preventivo, pestaña «El kit» y maqueta 3D (React Three Fiber).", "maqueta3d/src"],
        ["6. Salidas", "Borrador JSON, resumen por sector (CSV), aviso de acopio y difusión SMS SIMULADA, GLB de la maqueta y capturas.", "Descargas desde la aplicación"],
        ["7. Publicación", "Sitio estático en GitHub Pages; cada cambio en main ejecuta pruebas y compila antes de publicar.", ".github/workflows/pages.yml"],
    ], anchos=(3.0, 9.4, 4.2))

    h(doc, "4. Reglas explicables del prototipo")
    tabla(doc, ["Regla", "Criterio", "Responsable propuesto"], [
        ["Inundación", "Si el espacio cruza la amenaza de inundación del POT, se descarta.", "—"],
        ["Sismo", "Se excluyen cruces de licuación y corrimiento; el resto queda «requiere inspección antes de activar».", "—"],
        ["Sequía e incendios", "No se preseleccionan lugares sin evidencia suficiente.", "—"],
        ["Baños", "Redondeo superior de personas / 20 (Esfera).", "UAESP"],
        ["Agua", "Personas × 15 L al día (Esfera).", "EMCALI"],
        ["Área cubierta", "Personas × 3,5 m² (Esfera). La huella del predio no demuestra área cubierta útil.", "Gestión del Riesgo"],
        ["Dato desconocido", "Si no se conoce la existencia, la brecha queda «no calculable», nunca cero.", "—"],
    ], anchos=(3.0, 10.2, 3.4))
    par(doc, "La matriz de responsables es una propuesta del equipo basada en el Protocolo de Alojamientos Temporales de la UNGRD; falta validarla con las entidades.", "Nota: ")

    h(doc, "5. Tecnología y por qué es pertinente")
    tabla(doc, ["Elemento", "Elección", "Razón"], [
        ["Lenguaje", "TypeScript estricto", "Menos errores en reglas y cálculos"],
        ["Interfaz", "React + Vite", "Aplicación web ligera que corre en cualquier navegador"],
        ["Mapa", "Leaflet + OpenStreetMap", "Libre, sin clave de API ni costo por uso"],
        ["3D", "Three.js (React Three Fiber)", "Explica visualmente refugio, acopio y agua"],
        ["Pruebas", "Node test + Playwright", "22 pruebas unitarias de geometría, acopio, reglas y trazabilidad"],
        ["Alojamiento", "GitHub Pages", "Sitio estático público, sin servidor ni base de datos personal"],
    ], anchos=(3.0, 5.0, 8.6))

    h(doc, "6. Privacidad y seguridad desde el diseño")
    vinetas(doc, [
        "No hay registro nominal: solo conteos agregados; el código de registro de la maqueta es ficticio (DEMO-0001).",
        "No se piden ni guardan teléfonos; la difusión SMS es SIMULADA y no usa proveedor.",
        "La aplicación no solicita la geolocalización del usuario.",
        "Los datos exportados pasan por una lista explícita de atributos permitidos.",
        "Todo dato simulado aparece rotulado como SIMULADO en pantalla.",
    ])

    h(doc, "7. Arquitectura actual y arquitectura objetivo")
    tabla(doc, ["Aspecto", "Hoy (prototipo TRL 3)", "Objetivo (piloto)"], [
        ["Datos de los espacios", "Inventario abierto; disponibilidad «por confirmar» en todos los registros", "Ficha de caracterización en campo (Cali Lista) con superficie útil, accesos y servicios medidos"],
        ["Seguimiento", "Estados guardados en el navegador", "Base compartida con cuentas por entidad y bitácora de cambios"],
        ["Avisos a la comunidad", "SMS SIMULADO", "Canal oficial autorizado por la autoridad"],
        ["Retorno", "Representado en la maqueta", "Checklist y acta de entrega por espacio"],
        ["Entorno", "Cruce de amenazas", "Módulo preventivo: puntaje 60/40 del espacio y su red de apoyo"],
    ], anchos=(3.4, 6.4, 6.8))

    h(doc, "8. Vistas del prototipo")
    imagen(doc, CAPTURAS / "mapa.png", 15.5, "Mapa con los espacios públicos reales de Cali y sus cruces de amenaza.")
    imagen(doc, CAPTURAS / "ficha.png", 15.5, "Ficha de un espacio: datos, amenazas y evidencia pendiente.")
    imagen(doc, CAPTURAS / "exceso.png", 15.5, "Caso epou-9465 con reservas SIMULADAS: 20 módulos exceden el balance y se bloquean.")
    imagen(doc, CAPTURAS / "ajuste.png", 15.5, "Con 16 módulos la propuesta cumple el balance por superficie (no prueba encaje físico).")
    imagen(doc, RENDERS / "refugio_coliseo.png", 15.5, "Maqueta 3D ilustrativa del refugio: no representa construcciones existentes.")

    lineas_rojas(doc)
    fuentes(doc, ["epou", "comunas", "esfera", "osm", "ccby"], extra=[("Código del prototipo: ", URL_REPO)])
    pie(doc, "Territorio Preparado · Entregable 2 · Arquitectura de la solución")
    guardar(doc, "02_Arquitectura_de_la_solucion.docx")


# ---------- 3. Propuesta de valor ----------
def propuesta_valor():
    doc = nuevo_documento()
    portada(doc, 3, "Propuesta de valor", "Sprint de validación · Artefacto 4: propuesta de valor y plan (25 de septiembre)", "Valor potencial para ciudadanos y actores (20 %) · Innovación y diferenciación (15 %)")

    h(doc, "1. La propuesta en una frase")
    par(doc, "Para la Secretaría de Gestión del Riesgo de Emergencias y Desastres de Cali, que necesita decidir con criterios qué espacio público activar ante cada amenaza, qué le falta y quién debe responder, nuestra solución Territorio Preparado propone espacios candidatos y dimensiona su preparación mediante el cruce explicable de datos abiertos (inventario DAPM/IDESC, amenazas del POT y daños del sismo) con los estándares Esfera y la matriz de responsables de la UNGRD, a diferencia de hoy, cuando los espacios se habilitan sobre la marcha y la comunidad termina montando albergues por su cuenta en zonas inundables.")
    par(doc, "Territorio Preparado propone; la autoridad decide.", "Lema: ")

    h(doc, "2. Lienzo de valor")
    tabla(doc, ["Del lado del usuario", "Del lado de la solución"], [
        ["Tarea: decidir en las primeras horas qué espacio abrir ante una amenaza concreta.", "Preselección de hasta 3 candidatos con los descartes y criterios a la vista."],
        ["Dolor: no saber si un espacio sirve para esa amenaza (un buen refugio para sismo puede ser inundable).", "Reglas distintas por amenaza con capas oficiales del POT."],
        ["Dolor: no saber cuánto cabe sin sobreocupar.", "Balance de superficie con reservas explícitas y bloqueo por exceso."],
        ["Dolor: necesidades y responsables coordinados por teléfono.", "Necesidades Esfera por servicio y entidad responsable propuesta, con estado de seguimiento."],
        ["Dolor: cifras que no coinciden y ayudas con destino desconocido (Defensoría, OCHA).", "Borrador exportable con escenario, criterios, supuestos y pendientes."],
        ["Ganancia esperada: decisiones que se pueden explicar a la comunidad y a los entes de control.", "Cada propuesta dice por qué se propone un lugar, cuánto se puede ocupar y qué falta."],
    ], anchos=(8.3, 8.3))

    h(doc, "3. Valor por actor")
    tabla(doc, ["Actor", "Qué gana"], [
        ["Secretaría de Gestión del Riesgo", "Una propuesta de activación revisable y explicable en lugar de improvisar; conserva la decisión."],
        ["DATIC", "Una base técnica sobre datos abiertos, reproducible y sin datos personales."],
        ["EMCALI, UAESP, Salud, ICBF, Bienestar Social", "Saber qué servicio les corresponde en cada espacio y cuánto se necesita."],
        ["Familias y comunidades del oriente", "Avisos solo de puntos confirmados por la autoridad, sin ser identificadas; menos riesgo de refugiarse en zonas expuestas."],
        ["Usuarios habituales de parques y escenarios", "Un espacio que vuelve a su uso cotidiano con una ruta de retorno clara."],
    ], anchos=(5.2, 11.4))

    h(doc, "4. Qué nos diferencia")
    vinetas(doc, [
        ("Multiamenaza: ", "la aptitud de un espacio depende del tipo de amenaza. Ningún otro equipo del RETO-01 trabaja este enfoque."),
        ("Matriz de responsables: ", "cada necesidad se conecta con la entidad que debe cubrirla."),
        ("Albergues autogestionados: ", "incluimos a las comunidades que se refugian por su cuenta, un hallazgo propio que no está en la ficha del reto."),
        ("Honestidad con los datos: ", "un dato desconocido no vale cero; el sistema muestra lo que falta medir en vez de inventar capacidad."),
        ("Datos reales de Cali: ", "1.970 espacios públicos, amenazas del POT y daños satelitales del sismo, con fuentes y licencias."),
    ])
    tabla(doc, ["Otros equipos del RETO-01", "Enfoque"], [
        ["CityVe", "Condiciones sanitarias"],
        ["Certeza", "Agenda social y percepción de inseguridad"],
        ["Ecosistema Creativo", "Zonas afectadas y zonas seguras en tiempo real"],
        ["Revo HUD", "Motociclistas"],
        ["CityBeam", "Mobiliario urbano y paredes interactivas"],
        ["Nature Intelligence", "Alertas de incendios satelitales"],
        ["Aerostatic", "Comunicación entre comunidad e instituciones"],
    ], anchos=(5.2, 11.4))

    h(doc, "5. Un ejemplo que demuestra el valor")
    par(doc, "Espacio real epou-9465 (Parque · Colinas del Sur): huella cartográfica de unos 447,21 m². Con reservas SIMULADAS de 100 m² no utilizables, 80 m² de circulación, 60 m² de atención y 40 m² de otros usos, quedan 167,21 m². Con módulos SIMULADOS de 10 m², proponer 20 (200 m²) excede el balance y el sistema lo bloquea; 16 (160 m²) sí cumple.")
    par(doc, "Esto demuestra la detección de exceso por superficie. No demuestra que los módulos encajen físicamente, que quepan 16 familias ni que el parque esté autorizado.", "Límite: ")
    imagen(doc, CAPTURAS / "espacio-conceptual.png", 14.0, "Ilustración conceptual del espacio en modo emergencia. No es fotografía de Cali.")

    h(doc, "6. Supuestos críticos")
    tabla(doc, ["Supuesto", "Estado"], [
        ["Hay datos abiertos suficientes para una primera preselección de espacios por amenaza.", "Validado con evidencia"],
        ["La Secretaría usaría una recomendación explicable si la decisión de activar sigue siendo suya.", "Por validar"],
        ["EMCALI, UAESP, ICBF y Salud aceptan una matriz de responsables por servicio basada en el Protocolo de la UNGRD.", "Por validar"],
        ["Las juntas de acción comunal compartirían conteos agregados de los albergues autogestionados.", "Por validar"],
        ["Medir la superficie útil, los accesos y los servicios de los espacios priorizados es viable antes de la próxima temporada de lluvias.", "Por validar"],
    ], anchos=(12.6, 4.0))

    lineas_rojas(doc)
    fuentes(doc, ["defensoria", "ocha", "epou", "esfera"])
    pie(doc, "Territorio Preparado · Entregable 3 · Propuesta de valor")
    guardar(doc, "03_Propuesta_de_valor.docx")


# ---------- 4. Viabilidad GovTech ----------
def viabilidad():
    doc = nuevo_documento()
    portada(doc, 4, "Viabilidad GovTech", "Sprint de validación · Artefacto 4: nivel TRL/IRL y plan (25 de septiembre)", "Viabilidad técnica y operativa del concepto (10 %)")

    h(doc, "1. Encaje institucional")
    tabla(doc, ["Rol", "Entidad", "Qué hace con Territorio Preparado"], [
        ["Dueño funcional", "Secretaría de Gestión del Riesgo de Emergencias y Desastres", "Prepara escenarios, revisa las propuestas y decide la activación junto con el Consejo Municipal de Gestión del Riesgo."],
        ["Líder técnico", "DATIC", "Aloja la herramienta, integra los datos abiertos y administra accesos."],
        ["Fuente de datos", "DAPM/IDESC y Secretaría del Deporte", "Mantienen el inventario del espacio público y los escenarios."],
        ["Servicios", "EMCALI, UAESP, Salud, ICBF, Bienestar Social", "Reciben las brechas que les corresponden y reportan su estado."],
        ["Comunidad", "Juntas de acción comunal", "Reportan conteos agregados de albergues propios y reciben avisos confirmados."],
    ], anchos=(3.0, 5.2, 8.4))

    h(doc, "2. Encaje normativo")
    tabla(doc, ["Norma", "Cómo la cumple la solución"], [
        ["Ley 1523 de 2012 y Decreto 2157 de 2017", "La herramienta recomienda; la activación la decide la autoridad competente."],
        ["Ley 1712 de 2014 (transparencia)", "Usa datos abiertos y criterios que se pueden publicar."],
        ["Ley 1581 de 2012 y Ley 1098 de 2006", "Solo conteos agregados; ningún dato personal ni de menores."],
        ["Decreto 767 de 2022 (Gobierno Digital)", "Solución web, interoperable, con seguridad y privacidad desde el diseño."],
        ["Resolución MinTIC 1519 de 2020", "Contraste, lenguaje claro y controles accesibles."],
        ["Ley 400 de 1997, NSR-10 y Ley 1575 de 2012", "No reemplaza evaluaciones estructurales ni conceptos de Bomberos."],
        ["Circular SIC 002 de 2024 y CONPES 4144 de 2025 (verificar)", "Reglas explicables y decisión final humana."],
    ], anchos=(6.0, 10.6))

    h(doc, "3. Viabilidad técnica")
    vinetas(doc, [
        ("Ya funciona: ", "aplicación web pública en " + URL_DEMO + " con datos reales de Cali."),
        ("Código abierto y reproducible: ", "los datos se regeneran desde las fuentes originales con un comando; cada fuente guarda su URL, licencia y huella SHA-256."),
        ("Calidad comprobada: ", "compilación con TypeScript estricto y 22 pruebas automáticas que se ejecutan antes de cada publicación."),
        ("Sin dependencias costosas: ", "mapa libre (OpenStreetMap), sin claves de API; sitio estático sin servidor ni base de datos personal."),
        ("Portable: ", "puede alojarse en la infraestructura de DATIC o migrar a una base compartida (hoja de cálculo o base de datos) sin cambiar las reglas."),
    ])

    h(doc, "4. Viabilidad operativa")
    tabla(doc, ["Momento", "Quién", "Qué hace"], [
        ["Antes (preparación)", "DAPM, Deporte, Gestión del Riesgo", "Caracterizan los espacios priorizados y cargan superficie útil, accesos y servicios."],
        ["Alerta", "Gestión del Riesgo", "Define la amenaza y el escenario; obtiene candidatos y necesidades."],
        ["Decisión", "Secretaría o Consejo Municipal", "Revisa el borrador y decide qué se activa."],
        ["Operación", "Entidades de servicios", "Atienden las brechas asignadas y reportan su estado."],
        ["Aviso", "Autoridad", "Comunica solo los puntos confirmados por el canal oficial."],
        ["Retorno", "Gestión del Riesgo y administrador del espacio", "Checklist, retiro de equipos y acta de entrega."],
    ], anchos=(3.4, 4.8, 8.4))

    h(doc, "5. Costos: rubros identificados")
    par(doc, "No presentamos cifras de costos porque todavía no tenemos cotizaciones ni datos con fuente. Estos son los rubros que hay que estimar con la Alcaldía:")
    tabla(doc, ["Rubro", "Qué incluye", "Estado"], [
        ["Desarrollo y adaptación", "Cuentas por entidad, base compartida, acta de retorno", "Por estimar"],
        ["Caracterización en campo", "Medición de superficie útil, accesos y servicios de los espacios priorizados", "Por estimar"],
        ["Alojamiento y mantenimiento", "Infraestructura de DATIC; hoy el sitio estático no tiene costo", "Por estimar"],
        ["Canal de avisos", "Proveedor de mensajería autorizado y su cobertura", "Por estimar"],
        ["Kit modular (opcional)", "Módulos, cisternas y energía solar; compra previa por contratación ordinaria", "Por estimar"],
    ], anchos=(4.0, 9.0, 3.6))

    h(doc, "6. Riesgos y cómo los mitigamos")
    tabla(doc, ["Riesgo", "Mitigación"], [
        ["Que la herramienta se use como si decidiera", "Lenguaje de «candidatos», decisión explícita de la autoridad y borrador sin valor de acto administrativo."],
        ["Inventario desactualizado o sin capacidades reales", "Disponibilidad «por confirmar» por defecto; caracterización en campo antes de operar."],
        ["Tratamiento de datos personales", "Diseño sin datos personales: solo conteos agregados."],
        ["Resistencia de las entidades a la matriz de responsables", "Validar la matriz con cada entidad antes del piloto, basada en el Protocolo de la UNGRD."],
        ["Falta de conectividad en la emergencia", "Sitio ligero, geometrías locales y borrador descargable para trabajar sin conexión."],
    ], anchos=(6.0, 10.6))

    h(doc, "7. Ruta de adopción")
    tabla(doc, ["Fase", "Alcance", "Resultado esperado"], [
        ["0 · Hoy", "Prototipo con datos abiertos y escenarios SIMULADOS", "TRL 3 como meta: la lógica se demuestra en entorno controlado"],
        ["1 · Validación", "Sesiones con la Secretaría de Gestión del Riesgo y una JAC de Chiminangos o Calimio", "IRL 3 como meta: flujo de coordinación validado con actores reales"],
        ["2 · Piloto", "Espacios priorizados de las comunas del Jarillón (6, 7, 13, 14 y 21), antes de la temporada de lluvias", "Espacios caracterizados en campo y matriz de responsables acordada"],
        ["3 · Escalamiento", "Las 22 comunas de Cali, con el módulo preventivo y el acta de retorno", "Red de espacios preparados por amenaza"],
    ], anchos=(2.8, 7.4, 6.4))
    par(doc, "TRL 3 e IRL 3 son metas del hackathon, no certificaciones obtenidas.", "Aclaración: ")

    lineas_rojas(doc)
    fuentes(doc, ["epou", "esfera", "osm"], extra=[("Código del prototipo: ", URL_REPO)])
    pie(doc, "Territorio Preparado · Entregable 4 · Viabilidad GovTech")
    guardar(doc, "04_Viabilidad_GovTech.docx")


# ---------- 5. Elevator pitch ----------
PITCH = [
    ("0–15 s", "El gancho", "Mapa de la app filtrado por «Inundación» en Calimio",
     "El 10 de agosto, tras el sismo, 1.046 personas se refugiaron en canchas de Chiminangos y Calimio que la comunidad convirtió en albergues. Esas canchas están a entre 72 y 152 metros del dique del río Cauca: huyeron de un sismo hacia una zona inundable."),
    ("15–25 s", "El problema", "Cifra grande: «1.970 espacios públicos. ¿Cuál sirve para qué?»",
     "Cali tiene 1.970 espacios públicos inventariados, pero hoy nadie sabe cuál sirve ante cada amenaza, qué le falta ni quién debe responder."),
    ("25–48 s", "La solución", "Ficha del espacio y cálculo 20 → 16",
     "Territorio Preparado cruza datos abiertos del POT, IDESC y Copernicus. Eliges la amenaza y la población; el sistema descarta los espacios expuestos, propone candidatos, calcula cuánto cabe sin sobreocupar, cuánta agua y cuántos baños faltan, y qué entidad responde. Cada criterio queda a la vista."),
    ("48–60 s", "El cierre", "Logo y lema",
     "Ya funciona en línea con el mapa real de Cali. Que la próxima vez el parque esté listo antes de la emergencia. Territorio Preparado propone; la autoridad decide."),
]


def texto_pitch():
    return "\n\n".join(b[3] for b in PITCH)


def pitch():
    doc = nuevo_documento()
    portada(doc, 5, "Elevator pitch de 1 minuto (4.2)", "Presentación final de pitches · Pitch final y materiales de soporte", "Claridad y capacidad de síntesis del pitch (10 %)")
    palabras = len(texto_pitch().split())

    h(doc, "1. El texto completo")
    par(doc, f"{palabras} palabras, para leer a ritmo tranquilo en 60 segundos. Una sola voz. Si sobra tiempo, no se agrega nada: se habla más despacio.", "Duración: ")
    for tiempo, nombre, _, texto in PITCH:
        p = doc.add_paragraph()
        r = p.add_run(f"[{tiempo}] {nombre}. ")
        r.bold = True
        r.font.color.rgb = TITULO
        p.add_run(texto)

    h(doc, "2. Guion con diapositivas")
    tabla(doc, ["Tiempo", "Diapositiva", "Qué se dice"], [[t, f"{n}: {d}", x] for t, n, d, x in PITCH], anchos=(2.0, 4.6, 10.0))

    h(doc, "3. De dónde sale cada cifra")
    tabla(doc, ["Cifra", "Fuente"], [
        ["1.046 personas (336 familias) en 3 albergues autogestionados", "Defensoría del Pueblo"],
        ["A entre 72 y 152 m del dique; amenaza alta de inundación", "Cruce propio: amenazas del POT (IDESC, Acuerdo 0373 de 2014) + espacio público DAPM/IDESC"],
        ["1.970 espacios públicos inventariados", "DAPM/IDESC, datos abiertos de Cali"],
        ["20 → 16 módulos (si se muestra en pantalla)", "Cálculo por área del caso epou-9465 con reservas SIMULADAS"],
    ], anchos=(7.0, 9.6))

    h(doc, "4. Reglas para no perder puntos")
    vinetas(doc, [
        "Decir «candidatos que la autoridad revisa», nunca «los mejores lugares» ni «lugares habilitados» (Ley 1523 de 2012).",
        "El 20 → 16 es un cálculo por área con datos simulados: no afirmar que «caben 16 familias».",
        "No mencionar la cifra de fallecidos: las fuentes no coinciden (105 frente a 154).",
        "Decir «Territorio Preparado», que responde al reto Cali Activa.",
        "Ensayar con cronómetro al menos tres veces.",
    ])

    h(doc, "5. Respuestas de 10 segundos para el jurado")
    tabla(doc, ["Pregunta", "Respuesta"], [
        ["¿Qué es real y qué es simulado?", "Los espacios, las amenazas y los daños del sismo son reales y públicos. La población del escenario, las reservas de área, el IoT y los SMS están simulados y así se marcan en pantalla."],
        ["¿Y los datos personales?", "No se usan. Solo conteos agregados, sin datos de menores ni reconocimiento facial (Ley 1581 de 2012 y Ley 1098 de 2006)."],
        ["¿Reemplaza a los ingenieros o a Bomberos?", "No. Solo registra si existe una evaluación estructural vigente (NSR-10) y no emite conceptos de Bomberos (Ley 1575 de 2012)."],
        ["¿Lo validaron con usuarios?", "Tenemos evidencia documental de la Defensoría y OCHA. El siguiente paso es validarlo con la Secretaría de Gestión del Riesgo y las JAC de Chiminangos y Calimio."],
        ["¿Cómo escala?", "Es una aplicación web con datos abiertos. Con una base compartida, cada comuna puede alimentar su inventario."],
    ], anchos=(5.0, 11.6))

    fuentes(doc, ["defensoria", "epou"])
    pie(doc, "Territorio Preparado · Entregable 5 · Elevator pitch de 1 minuto (4.2)")
    guardar(doc, "05_Elevator_pitch_1_minuto.docx")
    return palabras


# ---------- 6. Bitácora de validación ----------
def bitacora():
    doc = nuevo_documento()
    portada(doc, 6, "Bitácora de validación", "Sprint de validación · Artefacto 5: prototipo, prueba y bitácora (26 de septiembre)", "Calidad de la validación de hipótesis (5 %) · Calidad y coherencia del prototipo (10 %)")

    h(doc, "1. Hipótesis y estado")
    tabla(doc, ["Hipótesis", "Cómo se valida", "Estado al 26 de septiembre"], [
        ["H1. Los espacios se activan sin saber si sirven para la amenaza.", "Evidencia documental y cruce de datos", "Apoyada: albergues autogestionados a 72–152 m del dique, en amenaza alta de inundación"],
        ["H2. Con datos abiertos se puede hacer una primera preselección por amenaza.", "Prototipo con 1.970 espacios y las capas del POT", "Validada técnicamente (TRL 3 como meta)"],
        ["H3. No está claro qué entidad responde por cada servicio.", "Defensoría, OCHA y entrevistas", "Apoyada por evidencia documental; entrevistas pendientes"],
        ["H4. La Secretaría usaría una recomendación explicable si conserva la decisión.", "Sesión con la Secretaría de Gestión del Riesgo", "Pendiente"],
        ["H5. Las JAC compartirían conteos agregados de sus albergues.", "Conversación con una JAC de Chiminangos o Calimio", "Pendiente"],
    ], anchos=(6.0, 4.8, 5.8))

    h(doc, "2. Evidencia de que el problema existe")
    vinetas(doc, [
        ("Defensoría del Pueblo: ", "336 familias (1.046 personas) en 3 albergues autogestionados; traslados de ayudas con destino desconocido por falta de caracterización."),
        ("OCHA: ", "cifras que no coinciden entre fuentes, familias sin información sobre inspecciones y fallas de agua y saneamiento en los albergues."),
        ("Semana: ", "al 31 de agosto, 22 de 43 escenarios deportivos seguían esperando inspección técnica."),
        ("Cruce propio de datos: ", "en las comunas del Jarillón, entre el 55 y el 92 % de los espacios públicos tiene amenaza alta de inundación."),
    ])

    h(doc, "3. Validación técnica del prototipo")
    tabla(doc, ["Prueba", "Resultado"], [
        ["Compilación de producción (TypeScript estricto)", "Correcta"],
        ["Pruebas automáticas (geometría, acopio, reglas de selección, trazabilidad sin datos inventados)", "22 de 22 aprobadas"],
        ["Pruebas de navegador de avisos, SMS y prevención en móvil", "Aprobadas el 25 de septiembre"],
        ["Comparación de 9 capas archivadas con los servicios WFS actuales", "Coincidieron en geometrías y atributos (25 de septiembre)"],
        ["Pestaña «El kit»: revisión a 1440 y 390 píxeles", "Tres imágenes con texto alternativo, sin desborde"],
        ["Publicación automática en GitHub Pages", "Pruebas y compilación antes de cada publicación"],
    ], anchos=(10.6, 6.0))
    par(doc, URL_DEMO, "Prototipo público: ")
    par(doc, "1.970 espacios públicos efectivos y 1.021 registros deportivos (DAPM/IDESC), amenazas del POT (inundación, licuación y remoción) y 1.090 puntos de daño del sismo (Copernicus EMSR916 e ICube-SERTIT). SIMULADOS y rotulados: la población del escenario, las reservas de área del acopio, el IoT del kit y la difusión SMS. Sin datos personales.", "Datos cargados: ")

    h(doc, "4. Rondas de prueba con usuarios")
    par(doc, "Las rondas con usuarios reales todavía no se han hecho. Se registran aquí cuando ocurran, solo con el rol de la persona, sin nombres ni datos de contacto. No se reemplazan con simulaciones.", "Estado: ")
    for n, rol in ((1, "Secretaría de Gestión del Riesgo (funcionario del área de preparación)"), (2, "Junta de acción comunal de Chiminangos o Calimio Norte"), (3, "Mentor del hackathon u otro equipo")):
        h(doc, f"Ronda {n}", 2)
        tabla(doc, None, [
            ["Con quién (rol)", f"PENDIENTE · previsto: {rol}"],
            ["Qué se mostró", "PENDIENTE"],
            ["Qué funcionó", "PENDIENTE"],
            ["Qué confundió o falló", "PENDIENTE"],
            ["Qué se ajustó", "PENDIENTE"],
        ], anchos=(4.6, 12.0), encabezado=False)

    h(doc, "5. Ajustes hechos durante el proceso")
    tabla(doc, ["Qué detectamos", "Qué ajustamos"], [
        ["«Peritaje» suena a dictamen técnico y cruza la línea roja de la NSR-10.", "Se habla de «caracterización» o «verificación de aptitud»."],
        ["Un parque en el mapa no es un lugar disponible.", "Todos los registros quedan «por confirmar»; un cruce sin amenaza no implica seguridad."],
        ["Proponer más módulos de los que caben sobreocupa el espacio.", "Balance de superficie con reservas explícitas y bloqueo por exceso (20 → 16)."],
        ["Una brecha desconocida mostrada como cero engaña.", "La brecha queda «no calculable» hasta tener la existencia real."],
        ["Para incendios no hay capa verificada.", "Los escenarios de incendio no preseleccionan candidatos sin evidencia."],
        ["Avisar sobre un punto no confirmado puede mover a la gente a un lugar inadecuado.", "El aviso indica «no acudir todavía» hasta que la autoridad confirme la apertura."],
        ["El botón Revisar no llevaba a la ficha seleccionada.", "Se corrigió la selección, el desplazamiento y el foco, también con teclado."],
        ["El nombre del reto se confundía con el del producto.", "Producto: Territorio Preparado; reto: Cali Activa."],
    ], anchos=(8.0, 8.6))

    h(doc, "6. Aprendizajes clave")
    vinetas(doc, [
        ("1. ", "Un espacio no es «apto» en general: su aptitud depende de la amenaza, y un buen refugio para el sismo puede estar en zona inundable."),
        ("2. ", "Un dato desconocido no vale cero: el sistema debe mostrar lo que falta medir en vez de inventar capacidad."),
        ("3. ", "Proponer no es habilitar: la comunidad solo debe recibir avisos de puntos que la autoridad ya confirmó."),
    ], estilo="Normal")

    h(doc, "7. Siguiente paso")
    par(doc, "Validar el flujo con la Secretaría de Gestión del Riesgo y una JAC del oriente; medir en campo la superficie útil y los servicios de los espacios priorizados; integrar el peritaje del entorno del módulo preventivo (puntaje 60/40) y un canal de avisos real y autorizado.")
    par(doc, "Con una semana más haríamos las tres rondas de prueba con usuarios y mediríamos cuánto tarda la Secretaría en generar una propuesta de activación con y sin la herramienta.", "Qué haríamos diferente: ")

    lineas_rojas(doc)
    fuentes(doc, ["defensoria", "ocha", "semana", "epou"], extra=[("Código y bitácora: ", URL_REPO)])
    pie(doc, "Territorio Preparado · Entregable 6 · Bitácora de validación")
    guardar(doc, "06_Bitacora_de_validacion.docx")


if __name__ == "__main__":
    alineacion()
    arquitectura()
    propuesta_valor()
    viabilidad()
    n = pitch()
    bitacora()
    (SALIDA / "src/pitch_4_2.txt").write_text(texto_pitch() + "\n", encoding="utf-8")
    print("Pitch 4.2:", n, "palabras")
