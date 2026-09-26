"""Cali Activa: presentación editorial editable, datos públicos y supuestos explícitos."""
from pathlib import Path
import json, math
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'entregables/presentacion'; ASSETS=OUT/'assets'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
DARK='102B30'; INK='173A3E'; TEAL='176E63'; MINT='BCEE7B'; CREAM='F4F2EB'; WHITE='FFFFFF'; GRAY='536D70'; LINE='D4DED5'; RED='AE4237'; PALE='E3ECD9'; SAND='EBD9B3'
FONT='DejaVu Sans'
SLIDES=[]

def color(h):return RGBColor.from_string(h)
def rect(s,x,y,w,h,fill,stroke=None,radius=False):
    sh=s.shapes.add_shape(SHAPE.ROUNDED_RECTANGLE if radius else SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if radius:sh.adjustments[0]=.09
    sh.fill.solid();sh.fill.fore_color.rgb=color(fill)
    if stroke:sh.line.color.rgb=color(stroke);sh.line.width=Pt(1)
    else:sh.line.fill.background()
    return sh

def text(s,x,y,w,h,content,size=20,fill=INK,bold=False,align=PP_ALIGN.LEFT):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
    tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,ln in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=ln
        p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(fill)
        p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(5);p.line_spacing=1.05
    return sh

def small(s,x,y,label,dark=False,w=10):return text(s,x,y,w,.28,label,9,'A8C6BE' if dark else GRAY)
def tag(s,x,y,label,dark=False,w=None):
    w=w or len(label)*.071+.32
    rect(s,x,y,w,.32,MINT if dark else PALE,radius=True)
    text(s,x+.12,y+.055,w-.2,.21,label,9,INK,True)
def circle(s,x,y,d,fill):
    sh=s.shapes.add_shape(SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=color(fill);sh.line.fill.background();return sh

def rule(s,x,y,w,fill=LINE):rect(s,x,y,w,.012,fill)
def arrow(s,x1,y1,x2,y2,fill=TEAL):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=color(fill);sh.line.width=Pt(1.7)
    tri=s.shapes.add_shape(SHAPE.ISOSCELES_TRIANGLE, Inches(x2-.055),Inches(y2-.055),Inches(.11),Inches(.11));tri.fill.solid();tri.fill.fore_color.rgb=color(fill);tri.line.fill.background();tri.rotation=90+math.degrees(math.atan2(y2-y1,x2-x1))
def node(s,x,y,w,h,title,fill=WHITE,stroke=LINE,size=17):
    sh=rect(s,x,y,w,h,fill,stroke,True);text(s,x+.16,y+.15,w-.32,h-.2,title,size,INK,True,PP_ALIGN.CENTER);return sh

def picture(s,path,x,y,w,h):
    im=Image.open(path);iw,ih=im.size;r=min(w/iw,h/ih);pw,ph=iw*r,ih*r
    return s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))

def new(label,title=None,dark=False,source=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color(DARK if dark else CREAM)
    i=len(prs.slides);fill=WHITE if dark else INK
    rect(s,.55,.4,.34,.34,MINT if dark else INK,radius=True);text(s,.595,.48,.26,.16,'CA',8,INK if dark else WHITE,True)
    text(s,1,.445,3,.2,'CALI ACTIVA',10,fill,True)
    text(s,8.1,.445,4.6,.2,label.upper(),9,'A8C6BE' if dark else GRAY,False,PP_ALIGN.RIGHT)
    if title:text(s,.6,1.13,12.1,1.3,title,32,fill,True)
    rule(s,.6,7.02,12.1,'315057' if dark else LINE)
    small(s,.6,7.15,source or 'RETO-01 · Cali Activa · Hackathon Smart City Expo Cali 2026',dark,w=11.2)
    text(s,12.08,7.12,.6,.25,f'{i:02}',10,fill,True,PP_ALIGN.RIGHT)
    return s

def notes(s,content,seconds=None):
    msg=(f'Tiempo sugerido: {seconds} segundos.\n\n' if seconds else '')+content
    s.notes_slide.notes_text_frame.text=msg
    SLIDES.append({'slide':len(prs.slides),'seconds':seconds,'notes':content})

# Cartography drawn from the project's actual public features, no fictional zones.
data=json.loads((ROOT/'maqueta3d/public/data/spaces.json').read_text())
communes=json.loads((ROOT/'maqueta3d/public/data/communes.json').read_text())
park=next(f for f in data['features'] if f['properties']['id']=='epou-9465')
public=[f for f in data['features'] if f['properties']['sourceKey']=='publicSpaces']
assert len(public)==1970

def rings(g):return g['coordinates'] if g['type']=='Polygon' else [r for p in g['coordinates'] for r in p]
def map_asset(name,bg,stroke,dot):
    w,h=1100,1350;im=Image.new('RGBA',(w,h),(0,0,0,0));d=ImageDraw.Draw(im)
    coords=[c for f in communes['features'] for r in rings(f['geometry']) for c in r]
    lo,hi=min(c[0] for c in coords),max(c[0] for c in coords);bot,top=min(c[1] for c in coords),max(c[1] for c in coords)
    scale=min((w-90)/(hi-lo),(h-90)/(top-bot));ox=(w-(hi-lo)*scale)/2;oy=(h-(top-bot)*scale)/2
    def p(c):return (ox+(c[0]-lo)*scale,oy+(top-c[1])*scale)
    for f in communes['features']:
        for r in rings(f['geometry']):
            pts=[p(c) for c in r];d.polygon(pts,fill=bg);d.line(pts+[pts[0]],fill=stroke,width=3)
    for f in public:
        x,y=p(f['properties']['coordinates']);d.ellipse((x-2,y-2,x+2,y+2),fill=dot)
    x,y=p(park['properties']['coordinates']);d.ellipse((x-13,y-13,x+13,y+13),fill='#E8B66A',outline='#FFFFFF',width=3)
    im.save(ASSETS/name)
map_asset('cali-oscuro.png','#173C40','#416E6B','#BCEE7B')
map_asset('cali-claro.png','#E0E9DB','#91AAA0','#176E63')

# 01: opening / large typographic treatment.
s=new('Pitch · espacios que cuidan',dark=True,source='Cartografía: Alcaldía de Santiago de Cali · DAPM / IDESC · adaptación CC BY-SA 4.0')
picture(s,ASSETS/'cali-oscuro.png',7.85,.91,4.75,5.98)
text(s,.6,1.55,7.2,.72,'CALI ACTIVA',45,WHITE,True)
text(s,.6,2.47,7.4,1.85,'Del espacio público\nal cuidado.',43,WHITE,True)
rect(s,.63,4.52,.85,.07,MINT)
text(s,.6,4.92,6.6,.84,'Preparar lugares. Coordinar ayudas.\nOrientar a la población.',21,'C4D7D0')
tag(s,.6,6.24,'DATOS REALES DE CALI',True,2.5);tag(s,3.28,6.24,'PROTOTIPO',True,1.55)
notes(s,'Cali Activa conecta la información del territorio con la preparación de espacios para emergencias. La pregunta central es cómo pasar de ver un lugar en el mapa a preparar una propuesta de atención que se pueda explicar y revisar.',15)

# 02: concrete pain, no unverifiable disaster statistics.
s=new('El problema','Un punto en el mapa\nno es un espacio listo.')
questions=[('01','¿Dónde?','Un lugar existente puede tener restricciones o no estar disponible.'),('02','¿Cuánto cabe?','La superficie total no revela el área libre para atender o almacenar.'),('03','¿Quién responde?','Las necesidades requieren seguimiento y responsables propuestos.')]
for j,(n,title,body) in enumerate(questions):
    x=.6+j*4.12;rule(s,x,3.08,3.77,TEAL);text(s,x,3.35,.9,.4,n,17,TEAL,True);text(s,x,4.02,3.8,.5,title,27,INK,True);text(s,x,4.8,3.62,1.15,body,18,GRAY)
text(s,.6,6.39,11.8,.4,'El reto es convertir información dispersa en una preparación trazable.',21,TEAL,True)
notes(s,'El inventario dice que un espacio existe. No confirma que esté abierto, que soporte la carga propuesta o que tenga agua y baños. El problema que abordamos es esa distancia entre localizar un espacio y preparar su uso, con variables y vacíos visibles.',20)

# 03: product's value chain.
s=new('La propuesta','Una ficha conecta\nterritorio, capacidad y coordinación.',dark=True)
labels=[('01','CONOCER','Espacio + amenaza'),('02','DIMENSIONAR','Superficie + carga'),('03','COORDINAR','Brechas + tareas'),('04','COMUNICAR','Aviso + condiciones'),('05','RECUPERAR','Revisión + retorno')]
for j,(n,a,b) in enumerate(labels):
    x=.63+j*2.49;circle(s,x,3.44,.49,MINT);text(s,x+.06,3.54,.36,.25,n,12,INK,True,PP_ALIGN.CENTER)
    if j<4:arrow(s,x+.62,3.68,x+2.25,3.68,'527A70')
    text(s,x,4.2,2.25,.36,a,14,MINT,True);text(s,x,4.75,2.18,.83,b,19,WHITE)
text(s,.65,6.24,12,.45,'El sistema recomienda. La autoridad decide la activación.',25,WHITE,True)
notes(s,'El producto une cinco momentos: conocer el lugar, dimensionar una propuesta, coordinar brechas, preparar comunicación y guiar el retorno. La recomendación no activa espacios. Esa decisión sigue en manos de la autoridad.',20)

# 04: public data evidence, not pretend capacity.
s=new('La base territorial','Cali real,\ncriterios visibles.',source='Fuente: Alcaldía de Santiago de Cali · DAPM / IDESC · corte 25 sep 2026 · CC BY-SA 4.0')
picture(s,ASSETS/'cali-claro.png',.7,2.46,4.1,4.36)
text(s,5.46,2.52,4.8,.88,'1.970',58,TEAL,True);text(s,5.52,3.42,6.6,.45,'registros de espacio público efectivo',20,INK,True)
for j,(a,b) in enumerate([('LOCALIZAR','Comuna, barrio y huella cartográfica.'),('FILTRAR','Amenaza, intersecciones y evidencia disponible.'),('EXPLICAR','Criterios, restricciones y datos por verificar.')]):
    y=4.19+j*.7;text(s,5.53,y,1.7,.27,a,12,TEAL,True);text(s,7.3,y-.015,5.13,.57,b,16,INK)
tag(s,5.53,6.42,'INVENTARIO ≠ DISPONIBILIDAD CONFIRMADA',w=5.17)
notes(s,'Trabajamos con 1.970 registros públicos de espacio público efectivo. El mapa representa geometrías reales y permite filtrar por territorio y amenaza. No sumamos los puntos deportivos como lugares únicos, porque puede haber duplicados. Tampoco convertimos inventario en disponibilidad: eso permanece por confirmar.',25)

# 05: explicit decision architecture.
s=new('La lógica','Tres decisiones antes de proponer un uso.')
columns=[(.6,'1 · AMENAZA','¿Pasa el filtro?','Con restricción o falta de evidencia: descartar o dejar pendiente.'),(4.77,'2 · SUPERFICIE','¿Cabe la carga?','Si excede el área restante: bloquear, ajustar y recalcular.'),(8.93,'3 · HABILITACIÓN','¿Está confirmado?','Sin confirmación: conservar borrador y no invitar a acudir.')]
for x,k,q,b in columns:
    tag(s,x,2.82,k,w=3.38)
    sh=s.shapes.add_shape(SHAPE.DIAMOND,Inches(x+.4),Inches(3.49),Inches(2.96),Inches(1.33));sh.fill.solid();sh.fill.fore_color.rgb=color(SAND);sh.line.fill.background()
    text(s,x+.7,3.97,2.36,.42,q,16,INK,True,PP_ALIGN.CENTER)
    arrow(s,x+1.88,4.82,x+1.88,5.16,RED)
    rect(s,x,5.3,3.8,1.05,WHITE,radius=True);text(s,x+.17,5.48,3.46,.82,b,16,INK)
text(s,.6,6.59,12,.22,'Pasar un filtro permite seguir revisando; no equivale a seguridad certificada.',13,GRAY)
notes(s,'El flujo tiene tres barreras claras. Primero, restricciones y evidencia por amenaza. Segundo, capacidad teórica por superficie. Tercero, habilitación de la autoridad. Un resultado favorable solo permite seguir con la propuesta. Ante incendios y sequía, sin evidencia suficiente, no habilitamos preselección. Google Maps ubica; no acredita una ruta segura.',30)

# 06: visually striking worked example with clear provenance.
s=new('Caso de uso · acopio','El tamaño del parque\nsí cambia la propuesta.',dark=True,source='epou-9465 · Colinas del Sur · área DAPM/IDESC calculada con Turf · reservas y módulos SIMULADOS')
tag(s,.65,2.69,'HUELLA REAL + SUPUESTOS SIMULADOS',True,4.52)
text(s,.65,3.37,2.6,.75,'447,21',42,WHITE,True);text(s,.67,4.17,2.58,.42,'m² cartográficos',16,'C4D7D0')
text(s,3.34,3.54,.5,.5,'−',30,MINT,True)
text(s,4.05,3.37,2.15,.75,'280',42,WHITE,True);text(s,4.09,4.17,2.63,.67,'m² reservados\nSIMULADOS',15,'C4D7D0')
rule(s,.65,4.99,6.15,'527A70')
text(s,.65,5.28,3.09,.66,'167,21 m²',32,MINT,True);text(s,3.93,5.33,2.88,.85,'para almacenamiento\n÷ 10 m² por módulo',16,WHITE)
text(s,.65,6.38,6.2,.4,'Reservas: 100 + 80 + 60 + 40 m²; ver desglose en notas.',10,'A8C6BE')
rect(s,7.32,2.64,5.35,3.98,'193D40',radius=True)
text(s,7.66,2.98,1.96,1.08,'16',65,MINT,True);text(s,9.72,3.25,2.5,.85,'módulos\nmáximo teórico',19,WHITE,True)
for j in range(20):rect(s,7.69+(j%10)*.44,4.38+(j//10)*.5,.32,.32,MINT if j<16 else 'D87060',radius=True)
text(s,7.67,5.64,4.51,.36,'20 exceden → ajustar a 16',20,WHITE,True)
text(s,7.67,6.17,4.5,.3,'Balance de área; no aforo ni encaje físico.',11,'C4D7D0')
notes(s,'Este parque real tiene 447,2149 metros cuadrados cartográficos. Para la demostración reservamos 100 no utilizables, 80 de circulación, 60 de atención y 40 de otros usos: 280 en total. Son supuestos simulados, no mediciones de campo. Quedan 167,21 metros cuadrados. A 10 por módulo, caben 16 módulos completos por balance de área. Veinte requieren 200 y exceden. Reducir a 16 deja 7,21 metros cuadrados. La forma del parque, accesos y carga aún requieren comprobación; no certificamos aforo ni encaje.',35)

# 07: coordination around needs.
s=new('Necesidades y coordinación','De contar personas\na preparar servicios.',source='Referencias de planificación: Manual Esfera 2018 · contextualizar; no certifican cumplimiento')
tag(s,.61,2.6,'ESCENARIO SIMULADO · 300 PERSONAS',w=4.53)
metrics=[('4.500','L de agua / día','300 × 15 L','EMCALI'),('15','baños','redondeo de 300 / 20','UAESP'),('1.050','m² cubiertos','300 × 3,5 m²','Gestión del Riesgo')]
for j,(v,u,f,who) in enumerate(metrics):
    x=.6+j*4.15;rect(s,x,3.25,3.81,2.54,WHITE,radius=True);text(s,x+.23,3.55,3.35,.74,v,44,TEAL,True);text(s,x+.25,4.43,3.27,.38,u,22,INK,True);text(s,x+.25,5.13,3.22,.3,f,13,GRAY);text(s,x,6.02,3.86,.31,who,16,INK,True)
text(s,.63,6.49,12,.28,'Responsables propuestos. Servicios existentes sin dato → brechas por medir.',16,TEAL,True)
notes(s,'Para un escenario simulado de 300 personas, mostramos las necesidades de referencia: 4.500 litros diarios, 15 baños y 1.050 metros cuadrados cubiertos. Son criterios de planificación a contextualizar, no servicios disponibles. Si no conocemos las existencias, la brecha queda por medir. Las entidades mostradas son responsables propuestos, sin validación con ellas en esta etapa.',25)

# 08: information reaches residents.
s=new('Prevención y comunicación','La población también necesita\nsaber cómo actuar.')
text(s,.63,2.78,6.4,.42,'ANTES: PREPARAR A LA COMUNIDAD',15,TEAL,True)
text(s,.63,3.42,6.05,1.53,'Aprender por amenaza.\nConsultar fuentes oficiales.\nConfirmar destino y apertura.',24,INK,True)
rule(s,.64,5.31,5.58)
text(s,.64,5.6,5.67,.86,'Encuentro, acopio y alojamiento cumplen funciones diferentes.',18,GRAY)
rect(s,7.37,2.62,4.5,3.99,DARK,radius=True);rect(s,7.59,2.87,4.06,3.48,WHITE,radius=True)
tag(s,7.82,3.08,'SMS SIMULADO',w=2.1)
text(s,7.82,3.72,3.56,.45,'Punto de acopio propuesto',17,INK,True)
text(s,7.82,4.35,3.46,.74,'Ubicación + horario\nInformación de ayudas',17,GRAY)
rect(s,7.8,5.32,3.63,.67,'F6E8D8',radius=True);text(s,8,5.48,3.2,.34,'Sin apertura: no acudir todavía.',13,RED,True)
text(s,7.5,6.65,4.7,.25,'A toda la población · cero envíos reales',12,TEAL,True)
notes(s,'La prevención es parte del producto: aprender por amenaza y distinguir encuentro, acopio y alojamiento. También preparamos mensajes generales de ubicación, horario y ayudas. En esta etapa el envío es simulado y no usamos teléfonos. Un aviso real que invite a acudir requiere apertura y condiciones confirmadas; el borrador no las inventa.',25)

# 09: evidence of an existing prototype, large true screenshot.
s=new('La demostración','El recorrido ya se puede probar.',source='Captura real del prototipo · datos IDESC y © OpenStreetMap · no confirma disponibilidad')
rect(s,.6,2.48,8.08,4.19,WHITE,radius=True)
picture(s,ASSETS/'mapa-app.png',.75,2.6,7.78,3.93)
steps=[('01','Seleccionar','Un espacio real del mapa.'),('02','Dimensionar','Reservas + carga propuesta.'),('03','Revisar','Brechas, borrador y SMS simulado.')]
for j,(n,title,body) in enumerate(steps):
    y=2.63+j*1.3;circle(s,9.15,y,.46,TEAL);text(s,9.19,y+.105,.38,.23,n,11,WHITE,True,PP_ALIGN.CENTER);text(s,9.82,y-.015,2.83,.38,title,22,INK,True);text(s,9.82,y+.47,2.77,.71,body,16,GRAY)
notes(s,'En la demostración seleccionamos un espacio real, abrimos Preparar este espacio, ingresamos las reservas y la carga de acopio, y observamos cómo se bloquea una propuesta que excede el área. Después revisamos necesidades, descargamos un borrador y mostramos la difusión simulada. No es una idea solo dibujada: este recorrido está implementado en el prototipo.',30)

# 10: close and team.
s=new('El valor público',dark=True)
text(s,.64,1.5,11.9,1.94,'Un espacio preparado.\nUna comunidad informada.',40,WHITE,True)
for j,(a,b) in enumerate([('EXPLICABLE','Cada propuesta muestra sus criterios.'),('COORDINADO','Necesidad, tarea y responsable propuesto.'),('REVERSIBLE','Revisión y retorno al uso cotidiano.')]):
    x=.65+j*4.17;rule(s,x,4.04,3.65,'527A70');text(s,x,4.33,3.8,.32,a,13,MINT,True);text(s,x,4.9,3.65,.85,b,18,WHITE)
text(s,.65,6.11,11.7,.43,'Siguiente paso: medir área útil, accesos y servicios de los espacios.',18,MINT,True)
small(s,.65,6.68,'Equipo: William Ortiz · Herlin Echeverry · Pablo Arango · Daniel Celis',True,w=12)
notes(s,'El valor es hacer explícitas las decisiones: por qué un lugar se revisa, cuánto se propone ocupar y qué falta resolver. El siguiente paso técnico es completar mediciones y servicios. Cali Activa propone y documenta; la autoridad decide. El equipo es William Ortiz, Herlin Echeverry, Pablo Arango y Daniel Celis.',15)

# 11: backup / honest readiness.
s=new('Anexo · alcance','Qué demuestra hoy\ny qué sigue por verificar.')
rect(s,.6,2.69,5.86,3.9,WHITE,radius=True);rect(s,6.78,2.69,5.89,3.9,WHITE,radius=True)
tag(s,.88,2.94,'IMPLEMENTADO',w=2.2);tag(s,7.05,2.94,'PENDIENTE / FUERA DE ESTA ETAPA',w=4.54)
for j,label in enumerate(['Inventario real y filtros por amenaza.','Balance de superficie y bloqueo por exceso.','Necesidades, brechas y exportación.','Educación y difusión SMS simulada.']):text(s,.91,3.69+j*.59,5.06,.54,label,17,INK)
for j,label in enumerate(['Área útil, accesos, aforos y servicios.','Encaje de módulos y condiciones de carga.','Disponibilidad y habilitación confirmadas.','Envío SMS real e integración operativa.']):text(s,7.06,3.69+j*.59,5.11,.54,label,17,INK)
notes(s,'Anexo para preguntas. No se afirma validación institucional, disponibilidad real, evaluación estructural ni certificación de seguridad contra incendios. La validación con entidades está fuera de esta etapa por decisión del equipo usuario. Los valores desconocidos siguen como desconocidos y los escenarios están marcados como simulados.')

# 12: source trail and attribution.
s=new('Anexo · fuentes','Datos, reglas y evidencia del prototipo.')
sources=[('01','ESPACIO PÚBLICO EFECTIVO','Alcaldía de Cali · DAPM / IDESC · CC BY-SA','https://datos.cali.gov.co/dataset/epou-espacio-publico-efectivo'),('02','COMUNAS DE SANTIAGO DE CALI','DAPM / IDESC · CC BY','https://www.datos.gov.co/dataset/Comunas-de-Santiago-de-Cali/dx2g-2mhm'),('03','REFERENCIAS HUMANITARIAS','Manual Esfera 2018 · agua, saneamiento y alojamiento','https://spherestandards.org/wp-content/uploads/Sphere-Handbook-2018-EN.pdf'),('04','BITÁCORA Y CATÁLOGO DEL PROYECTO','docs/PROYECTO.md · prototipo/datos/FUENTES.md','https://github.com/leonidas452528/hackaton-camaracomercio2026')]
for j,(n,title,desc,url) in enumerate(sources):
    y=2.6+j*.91;text(s,.63,y,.5,.3,n,12,TEAL,True);text(s,1.29,y,10.6,.3,title,14,INK,True);text(s,1.29,y+.34,10.7,.27,desc,12,GRAY)
    link=text(s,1.29,y+.62,10.8,.23,url,9,TEAL);link.text_frame.paragraphs[0].runs[0].hyperlink.address=url
text(s,.63,6.54,12,.25,'Corte del inventario: 25 sep 2026 · Adaptaciones cartográficas: CC BY-SA 4.0 · Sin datos personales.',12,GRAY)
notes(s,'Fuentes de la presentación. Los gráficos cartográficos derivan de los archivos públicos ya archivados en el proyecto; no se agregó un dataset. El caso de acopio usa el polígono de epou-9465 y la misma función de cálculo del frontend. La captura es del prototipo local. Las reservas, el tamaño de módulo y la población del escenario son simulados, marcados como tales.')

prs.core_properties.title='Cali Activa — Del espacio público al cuidado'
prs.core_properties.subject='Pitch del prototipo multiamenaza de Cali Activa'
prs.core_properties.author='Equipo Cali Activa'
prs.core_properties.keywords='Cali, espacio público, acopio, prevención, prototipo'
prs.save(OUT/'Cali_Activa_Presentacion.pptx')
(OUT/'src/guion.json').write_text(json.dumps(SLIDES,ensure_ascii=False,indent=2)+'\n')
(OUT/'Guion_del_pitch.md').write_text('# Guion del pitch · Cali Activa\n\n10 diapositivas principales · 4 minutos sugeridos. Las diapositivas 11–12 son anexos.\n\n'+'\n\n'.join(f'## Diapositiva {r["slide"]}'+(f' · {r["seconds"]} segundos' if r['seconds'] else ' · Anexo')+'\n\n'+r['notes'] for r in SLIDES)+'\n')
print('Presentación creada:',len(prs.slides),'diapositivas; pitch',sum(r['seconds'] or 0 for r in SLIDES),'segundos.')
