from pathlib import Path
import json,re
import pymupdf as fitz
from pptx import Presentation
from PIL import Image,ImageOps,ImageDraw
ROOT=Path(__file__).resolve().parents[1]
doc=fitz.open(ROOT/'Cali_Activa_Presentacion.pdf');prs=Presentation(ROOT/'Cali_Activa_Presentacion.pptx')
assert len(doc)==len(prs.slides)==12
problems=[]
norm=lambda s: re.sub(r'\s+','',s)
for i,(page,slide) in enumerate(zip(doc,prs.slides)):
    actual=norm(page.get_text())
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text.strip():
            if norm(shape.text) not in actual:problems.append({'slide':i+1,'text':shape.text})
        if shape.left<0 or shape.top<0 or shape.left+shape.width>prs.slide_width+500 or shape.top+shape.height>prs.slide_height+500:
            problems.append({'slide':i+1,'outsideCanvas':shape.name})
    assert abs(page.rect.width/page.rect.height-16/9)<.001
    assert slide.has_notes_slide
    pix=page.get_pixmap(matrix=fitz.Matrix(1.33333,1.33333),alpha=False)
    pix.save(str(ROOT/f'assets/slide-{i+1:02}.png'))
for half in range(2):
    sheet=Image.new('RGB',(1280,1080),'#D5DEDA')
    for k in range(6):
        im=Image.open(ROOT/f'assets/slide-{half*6+k+1:02}.png').resize((632,355))
        sheet.paste(im,((k%2)*640,(k//2)*360))
    sheet.save(ROOT/f'Vista_general_{half+1}.jpg',quality=94)
Image.open(ROOT/'assets/slide-01.png').save(ROOT/'Portada.jpg',quality=95)
result={'slides':12,'mainSlides':10,'appendixSlides':2,'aspectRatio':'16:9','speakerNotes':12,'mainPitchSeconds':sum(r['seconds'] or 0 for r in json.loads((ROOT/'src/guion.json').read_text())),'editableTextBoxes':sum(1 for s in prs.slides for sh in s.shapes if sh.has_text_frame and sh.text.strip()),'textOrBoundsProblems':problems,'pdfBytes':(ROOT/'Cali_Activa_Presentacion.pdf').stat().st_size,'pptxBytes':(ROOT/'Cali_Activa_Presentacion.pptx').stat().st_size}
(ROOT/'verificacion.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(result,ensure_ascii=False,indent=2))
