"""Territorio Preparado: presentación editorial editable, datos públicos y supuestos explícitos."""
from pathlib import Path
import json, math
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'entregables/presentacion-ampliada'; ASSETS=OUT/'assets'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
DARK='0C2530'; INK='14323D'; TEAL='087F80'; MINT='6AE9C7'; CREAM='F4F7F8'; WHITE='FFFFFF'; GRAY='536D70'; LINE='D4DED5'; RED='AE4237'; PALE='E3ECD9'; SAND='EBD9B3'
FONT='DejaVu Sans'
BRAND=json.loads((OUT/'src/marca.json').read_text())
SLIDES=[]

def color(h):return RGBColor.from_string(h)
def rect(s,x,y,w,h,fill,stroke=None,radius=False):
    sh=s.shapes.add_shape(SHAPE.ROUNDED_RECTANGLE if radius else SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if radius:sh.adjustments[0]=.09
    sh.fill.solid();sh.fill.fore_color.rgb=color(fill)
    if stroke:sh.line.color.rgb=color(stroke);sh.line.width=Pt(1)
    else:sh.line.fill.background()
    return sh

def text(s,x,y,w,h,content,size=20,fill=INK,bold=False,align=PP_ALIGN.LEFT):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
    tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,ln in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=ln
        p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(fill)
        p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(5);p.line_spacing=1.05
    return sh

def small(s,x,y,label,dark=False,w=10):return text(s,x,y,w,.28,label,9,'A8C6BE' if dark else GRAY)
def tag(s,x,y,label,dark=False,w=None):
    w=w or len(label)*.071+.32
    rect(s,x,y,w,.32,MINT if dark else PALE,radius=True)
    text(s,x+.12,y+.055,w-.2,.21,label,9,INK,True)
def circle(s,x,y,d,fill):
    sh=s.shapes.add_shape(SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=color(fill);sh.line.fill.background();return sh

def rule(s,x,y,w,fill=LINE):rect(s,x,y,w,.012,fill)
def arrow(s,x1,y1,x2,y2,fill=TEAL):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=color(fill);sh.line.width=Pt(1.7)
    tri=s.shapes.add_shape(SHAPE.ISOSCELES_TRIANGLE, Inches(x2-.055),Inches(y2-.055),Inches(.11),Inches(.11));tri.fill.solid();tri.fill.fore_color.rgb=color(fill);tri.line.fill.background();tri.rotation=90+math.degrees(math.atan2(y2-y1,x2-x1))
def node(s,x,y,w,h,title,fill=WHITE,stroke=LINE,size=17):
    sh=rect(s,x,y,w,h,fill,stroke,True);text(s,x+.16,y+.15,w-.32,h-.2,title,size,INK,True,PP_ALIGN.CENTER);return sh

def picture(s,path,x,y,w,h):
    im=Image.open(path);iw,ih=im.size;r=min(w/iw,h/ih);pw,ph=iw*r,ih*r
    return s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))

def new(label,title=None,dark=False,source=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color(DARK if dark else CREAM)
    i=len(prs.slides);fill=WHITE if dark else INK
    rect(s,.55,.4,.34,.34,MINT if dark else INK,radius=True);text(s,.595,.48,.26,.16,BRAND['initials'],8,INK if dark else WHITE,True)
    text(s,1,.445,3,.2,BRAND['name'].upper(),10,fill,True)
    text(s,8.1,.445,4.6,.2,label.upper(),9,'A8C6BE' if dark else GRAY,False,PP_ALIGN.RIGHT)
    if title:text(s,.6,1.13,12.1,1.3,title,32,fill,True)
    rule(s,.6,7.02,12.1,'315057' if dark else LINE)
    small(s,.6,7.15,source or 'Producto: '+BRAND['name']+' · Responde al RETO-01 Cali Activa · Prototipo',dark,w=11.2)
    text(s,12.08,7.12,.6,.25,f'{i:02}',10,fill,True,PP_ALIGN.RIGHT)
    return s

def notes(s,content,seconds=None):
    msg=(f'Tiempo sugerido: {seconds} segundos.\n\n' if seconds else '')+content
    s.notes_slide.notes_text_frame.text=msg
    SLIDES.append({'slide':len(prs.slides),'seconds':seconds,'notes':content})

# Cartography drawn from the project's actual public features, no fictional zones.
data=json.loads((ROOT/'maqueta3d/public/data/spaces.json').read_text())
communes=json.loads((ROOT/'maqueta3d/public/data/communes.json').read_text())
park=next(f for f in data['features'] if f['properties']['id']=='epou-9465')
public=[f for f in data['features'] if f['properties']['sourceKey']=='publicSpaces']
assert len(public)==1970

def rings(g):return g['coordinates'] if g['type']=='Polygon' else [r for p in g['coordinates'] for r in p]
def map_asset(name,bg,stroke,dot):
    w,h=1100,1350;im=Image.new('RGBA',(w,h),(0,0,0,0));d=ImageDraw.Draw(im)
    coords=[c for f in communes['features'] for r in rings(f['geometry']) for c in r]
    lo,hi=min(c[0] for c in coords),max(c[0] for c in coords);bot,top=min(c[1] for c in coords),max(c[1] for c in coords)
    scale=min((w-90)/(hi-lo),(h-90)/(top-bot));ox=(w-(hi-lo)*scale)/2;oy=(h-(top-bot)*scale)/2
    def p(c):return (ox+(c[0]-lo)*scale,oy+(top-c[1])*scale)
    for f in communes['features']:
        for r in rings(f['geometry']):
            pts=[p(c) for c in r];d.polygon(pts,fill=bg);d.line(pts+[pts[0]],fill=stroke,width=3)
    for f in public:
        x,y=p(f['properties']['coordinates']);d.ellipse((x-2,y-2,x+2,y+2),fill=dot)
    x,y=p(park['properties']['coordinates']);d.ellipse((x-13,y-13,x+13,y+13),fill='#E8B66A',outline='#FFFFFF',width=3)
    im.save(ASSETS/name)
map_asset('cali-oscuro.png','#173C40','#416E6B','#BCEE7B')
map_asset('cali-claro.png','#E0E9DB','#91AAA0','#176E63')


def table(s,headers,rows,widths,y=2.5,x=.6,rowh=.65):
    xx=x
    for head,w in zip(headers,widths):
        rect(s,xx,y,w,.46,DARK);text(s,xx+.12,y+.11,w-.24,.27,head,11,WHITE,True);xx+=w
    for i,row in enumerate(rows):
        xx=x;yy=y+.46+i*rowh
        for j,(cell,w) in enumerate(zip(row,widths)):
            rect(s,xx,yy,w,rowh,WHITE if i%2==0 else 'E7EFF1')
            text(s,xx+.12,yy+.14,w-.24,rowh-.18,cell,13,INK,j==0);xx+=w

def paragraph_label(s,x,y,n,title,body,w=3.65):
    circle(s,x,y,.39,TEAL);text(s,x+.06,y+.07,.27,.22,n,11,WHITE,True,PP_ALIGN.CENTER)
    text(s,x+.57,y-.01,w-.57,.5,title,18,INK,True);text(s,x+.57,y+.64,w-.57,1.1,body,15,GRAY)

# 01 — clear product identity and a conceptual scene, not an operational claim.
s=new('Preparación de espacios + coordinación de ayudas',dark=True,source='Ilustración conceptual generada con IA · no representa un parque real, inventario ni diseño aprobado')
picture(s,ASSETS/'espacio-conceptual.png',5.87,1.37,7.0,4.15)
text(s,.62,1.45,5.15,1.65,BRAND['name'].replace(' ', '\n', 1),42,WHITE,True)
text(s,.65,3.35,5.05,1.2,BRAND['tagline'],24,WHITE,True)
text(s,.65,4.76,5.15,1,'Dónde preparar.\nCuánto proponer.\nQué falta resolver.',21,'C1DADE')
for j,label in enumerate(['TERRITORIO','CAPACIDAD','COMUNIDAD']):tag(s,6.04+j*2.19,5.95,label,True,2.02)
tag(s,.65,6.34,'PLANEACIÓN PARA EMERGENCIAS',True,4.7)
notes(s,BRAND['name']+' expresa el propósito: preparar espacios para emergencias. El sistema conecta ubicación, capacidad propuesta, servicios y comunicación. La imagen es conceptual: no acredita que un parque esté construido o habilitado así.',10)

# 02 — a concrete trigger and the three outputs, no fabricated impacts.
s=new('La situación que debemos resolver','Hay personas por atender.\n¿Está listo el lugar?')
tag(s,.65,2.55,'CASO DE DEMOSTRACIÓN · POBLACIÓN SIMULADA',w=6.1)
text(s,.66,3.26,3.35,1,'300',69,TEAL,True);text(s,.67,4.41,3.25,.67,'personas en el\nescenario de atención',19,INK,True)
rule(s,4.39,2.87,.02)
for j,(q,b) in enumerate([('UBICACIÓN','¿Qué espacios tienen restricciones por la amenaza?'),('OCUPACIÓN','¿Qué área queda para acopio, circulación y atención?'),('COORDINACIÓN','¿Qué servicios faltan y a quién se propone dar seguimiento?')]):
    y=2.97+j*1.05;text(s,4.79,y,7.24,.34,q,13,TEAL,True);text(s,4.79,y+.42,7.4,.55,b,19,INK)
text(s,.66,6.46,12,.38,'La salida es una propuesta revisable, con criterios y vacíos visibles.',20,TEAL,True)
notes(s,'Tomemos un escenario simulado de 300 personas. No basta elegir el parque más cercano. Debemos revisar amenaza, espacio disponible y servicios. El producto prepara una propuesta que se puede revisar; no toma la decisión de activación.',15)

# 03 — role clarity, separate process from implemented access control.
s=new('Usuarios y responsabilidades','Tres actores. Una información conectada.')
for j,(title,verb,body,out) in enumerate([
('EQUIPO DE PREPARACIÓN','Caracteriza y propone','Consulta el territorio, registra supuestos, calcula necesidades y revisa brechas.','Entrega: borrador explicable'),
('AUTORIDAD COMPETENTE','Revisa y decide','Determina habilitación y condiciones de apertura. La aplicación no sustituye esa decisión.','Salida externa: decisión de activación'),
('COMUNIDAD','Se prepara y se informa','Aprende por amenaza, consulta avisos oficiales y distingue dónde aportar o recibir información.','Canal de demo: SMS simulado')]):
    x=.6+j*4.16;rect(s,x,2.74,3.82,3.83,WHITE,radius=True);tag(s,x+.19,2.94,title,w=3.44)
    text(s,x+.2,3.57,3.43,.94,verb,25,INK,True);text(s,x+.2,4.62,3.4,1.03,body,16,GRAY);rule(s,x+.2,5.89,3.41);text(s,x+.2,6.07,3.4,.37,out,12,TEAL,True)
notes(s,'El equipo prepara; la autoridad decide; la comunidad aprende y se informa. Son roles del proceso propuesto, no una afirmación de que ya existe integración con las entidades ni un sistema de permisos operativos validado.',15)

# 04 — evidence-rich inventory and a real record.
s=new('Evidencia territorial','El recorrido empieza en datos públicos de Cali.',source='Alcaldía de Santiago de Cali · DAPM / IDESC · corte 25 sep 2026 · adaptaciones CC BY-SA 4.0')
picture(s,ASSETS/'cali-claro.png',.6,2.42,3.53,4.3)
text(s,4.46,2.42,3.49,.85,'1.970',54,TEAL,True);text(s,4.47,3.34,7.6,.42,'registros de espacio público efectivo',20,INK,True)
node(s,4.47,4.12,3.71,1.17,'Geometría + comuna\n+ barrio + fuente',size=18)
node(s,8.49,4.12,4.09,1.17,'Cruces por amenaza\n+ restricciones + vacíos',size=18)
rect(s,4.47,5.58,8.13,.91,DARK,radius=True);text(s,4.69,5.76,7.65,.54,'Caso real: epou-9465 · Parque Colinas del Sur\n447,21 m² cartográficos · apertura no confirmada',16,WHITE,True)
notes(s,'La base contiene 1.970 registros de espacio público efectivo. Conservamos fuente, geometría y cruces por amenaza. No sumamos puntos deportivos duplicables como lugares únicos. El caso que vamos a seguir es epou-9465, en Colinas del Sur: su existencia y área cartográfica no prueban disponibilidad.',20)

# 05 — explicit threat-dependent behavior.
s=new('Lógica multiamenaza','Cambiar la amenaza cambia la revisión.')
table(s,['AMENAZA','CRITERIO ACTUAL','SALIDA DEL PROTOTIPO'],[
('Inundación','Cruce con capas de inundación.','Con cruce: descartado. Sin cruce: condicionado.'),
('Sismo','Cruce de licuación / corrimiento.','Con cruce: fuera de preselección. Sin cruce: requiere inspección.'),
('Incendio forestal','Sin evidencia suficiente del incidente y restricciones.','No se preselecciona; tareas de preparación por revisar.'),
('Incendio en edificación','Faltan condiciones del incidente y evaluaciones.','No se preselecciona; no se emiten conceptos de Bomberos.'),
('Sequía','Sin datos suficientes de sequía y abastecimiento.','No se recomienda un espacio por este escenario.')
],[2.17,4.39,5.54],y=2.53,rowh=.67)
text(s,.64,6.53,12,.35,'Sin cruce detectado ≠ ausencia de riesgo. La evidencia insuficiente no se convierte en aptitud.',15,TEAL,True)
notes(s,'La aptitud depende de la amenaza. Inundación y sismo usan cruces públicos disponibles. En incendios y sequía faltan evidencias para preseleccionar: el prototipo no inventa esa aptitud. Para sismo sigue siendo necesaria la evaluación correspondiente; no se certifican estructuras ni seguridad contra incendios.',25)

# 06 — detailed decision flow, not a decorative line of stages.
s=new('Flujo de decisión','Cada avance tiene una condición.',dark=True)
xs=[.66,3.84,7.02,10.2]
for x,title in zip(xs,['Definir escenario','Consultar candidatos','Dimensionar acopio','Preparar borrador']):node(s,x,2.82,2.49,.9,title,fill='E5F2F0',stroke=TEAL,size=17)
for a,b in zip(xs,xs[1:]):arrow(s,a+2.49,3.27,b,3.27,MINT)
for j,(x,q,no) in enumerate([(.66,'¿Evidencia suficiente?','NO → dejar pendiente'),(3.84,'¿Sin restricción?','NO → revisar otro espacio'),(7.02,'¿Cabe por superficie?','NO → ajustar y recalcular'),(10.2,'¿Apertura confirmada?','NO → no invitar a acudir')]):
    arrow(s,x+1.25,3.72,x+1.25,4.07,MINT)
    rect(s,x,4.2,2.49,.94,'163D46',radius=True);text(s,x+.13,4.43,2.23,.63,q,17,WHITE,True,PP_ALIGN.CENTER)
    text(s,x,5.48,2.49,.78,no,15,'EAB395',True)
text(s,.67,6.43,11.98,.32,'Con datos y condiciones favorables se continúa; activar sigue siendo decisión de la autoridad.',17,WHITE)
notes(s,'El flujo permite detenerse y volver. Sin evidencia, queda pendiente; con restricciones, se revisa otro espacio; con sobreocupación, se recalcula. Incluso un borrador completo no invita a acudir hasta que la apertura esté confirmada. La preselección considera tres candidatos cercanos por distancia en línea recta, no rutas seguras.',25)

# 07 — meaningful conceptual visual, clearly labeled.
s=new('Organización del espacio','El acopio convive con atención, acceso y servicios.',source='Ilustración conceptual generada con IA · no es implantación automática, parque real ni diseño técnico aprobado')
picture(s,ASSETS/'espacio-conceptual.png',.55,2.48,8.63,4.59)
for x,y,n in [(2.19,3.26,'1'),(4.34,4.25,'2'),(5.4,3.43,'3'),(6.84,5.61,'4')]:
    circle(s,x,y,.37,DARK);text(s,x+.055,y+.072,.26,.22,n,11,WHITE,True,PP_ALIGN.CENTER)
for j,(n,a,b) in enumerate([('1','ALMACENAMIENTO','Ayudas protegidas y ordenadas.'),('2','ATENCIÓN','Recepción, entrega y espera.'),('3','SERVICIOS','Agua y operación por verificar.'),('4','CIRCULACIÓN','Accesos libres y rutas revisables.')]):
    y=2.58+j*.99;circle(s,9.47,y,.3,TEAL);text(s,9.51,y+.055,.23,.19,n,9,WHITE,True,PP_ALIGN.CENTER);text(s,9.94,y,2.76,.28,a,12,TEAL,True);text(s,9.94,y+.39,2.72,.57,b,15,INK)
notes(s,'Esta imagen conceptual muestra por qué no todo el predio es almacenamiento. Se necesitan áreas diferenciadas de atención, circulación y servicios. No es la distribución del parque del caso ni una obra diseñada por el sistema. Las cubiertas y el tanque ilustran funciones, no inventario ni capacidad verificadas.',20)

# 08 — full allocation breakdown, exact provenance.
s=new('Caso real + supuestos','De 447,21 m² a 167,21 m² para proponer acopio.',source='Área real epou-9465: 447,2149007 m² · reservas SIMULADAS · cálculo del frontend, sin aforo certificado')
text(s,.63,2.7,5.54,.43,'¿Qué se descuenta?',23,INK,True)
items=[('Área no utilizable',100,'879BA2'),('Circulación y accesos',80,'B7C8CA'),('Atención y espera',60,'F1BB73'),('Otros usos / ocupación',40,'C5D8EB'),('Almacenamiento restante',167.2149,TEAL)]
x=.65
for label,value,c in items:
    w=12.0*value/447.2149;rect(s,x,3.43,w,.94,c);text(s,x+.04,3.7,w-.08,.32,('167,21' if value>150 else str(int(value))),18,WHITE if c==TEAL else INK,True,PP_ALIGN.CENTER);x+=w
for j,(label,value,c) in enumerate(items):
    x=.66+j*2.43;circle(s,x,4.72,.15,c);text(s,x+.25,4.65,2.04,.84,label,14,INK,True)
rect(s,.65,5.8,12,.71,DARK,radius=True);text(s,.9,5.97,11.45,.4,'447,21 − (100 + 80 + 60 + 40) = 167,21 m² disponibles en el supuesto',20,WHITE,True)
notes(s,'Usamos la huella real de Colinas del Sur. Los descuentos son un ejemplo simulado: 100 metros no utilizables, 80 de circulación, 60 de atención y 40 de otros usos. No son mediciones. El resultado es 167,21 metros cuadrados. Un campo sin dato no vale cero; sin polígono o con datos incompletos, no se calcula capacidad utilizable.',30)

# 09 — actual application outputs, before and after.
s=new('Demostración del control','20 exceden. 16 pasan el balance por área.')
for x,num,title,detail,c in [(.64,'20','PROPUESTA BLOQUEADA','200 m² pedidos → excede 32,79 m²',RED),(6.98,'16','PROPUESTA AJUSTADA','160 m² pedidos → restan 7,21 m²',TEAL)]:
    rect(s,x,2.65,5.7,1.68,WHITE,radius=True);text(s,x+.2,2.88,1.4,.9,num,56,c,True);text(s,x+1.83,2.95,3.59,.6,title,15,c,True);text(s,x+1.83,3.56,3.61,.52,detail,14,INK)
picture(s,ASSETS/'exceso.png',.65,4.67,5.7,1.18);picture(s,ASSETS/'ajuste.png',6.98,4.67,5.7,1.18)
text(s,.65,6.1,11.96,.34,'Módulo supuesto: 10 m². Máximo: parte entera de 167,21 / 10 = 16.',18,TEAL,True)
text(s,.65,6.62,11.96,.22,'Capturas reales del prototipo · editar medidas o cambiar espacio invalida la inclusión anterior.',12,GRAY)
notes(s,'Comprobamos la regla en la aplicación. Veinte módulos de diez metros cuadrados requieren doscientos y no caben en el balance. Dieciséis requieren ciento sesenta, por lo que se puede incluir la distribución en el borrador. La capacidad operacional sigue sin verificar. No se demuestra encaje en la forma irregular del predio ni aforo de personas.',25)

# 10 — full service gap trace.
s=new('Coordinación de necesidades','Cada necesidad tiene una brecha y un seguimiento.',source='Referencias Esfera 2018: contextualizar los valores · responsables propuestos, sin validación con entidades')
tag(s,.64,2.55,'ESCENARIO SIMULADO · 300 PERSONAS',w=4.9)
table(s,['SERVICIO','NECESIDAD','EXISTENCIA','BRECHA','RESPONSABLE PROPUESTO'],[
('Agua básica','4.500 L / día','Sin dato','Por medir','EMCALI'),
('Baños','15 unidades','Sin dato','Por medir','UAESP'),
('Área cubierta','1.050 m²','Sin dato','Por medir','Gestión del Riesgo')
],[2.1,2.12,1.9,1.85,4.13],y=3.2,rowh=.69)
text(s,.66,6.04,12,.63,'Necesidad calculada ≠ servicio disponible.\nEl seguimiento local distingue “Por medir” y “En revisión”.',20,TEAL,True)
notes(s,'Las referencias para 300 personas producen estas necesidades. Como no sabemos qué existe, no restamos cero ni afirmamos faltantes reales. Se conserva la brecha por medir y una entidad propuesta. El seguimiento actual es local: no representa asignación aceptada por una entidad externa.',20)

# 11 — actual prevention and simulated delivery outputs.
s=new('Preparación ciudadana + difusión','Orientar antes. Informar con condiciones claras.')
text(s,.65,2.62,5.8,.4,'CULTURA PREVENTIVA',16,TEAL,True)
picture(s,ASSETS/'prevencion.png',.63,3.22,6.17,1.23)
text(s,.65,4.68,5.94,.88,'¿Aparecer en el mapa significa\nque puedo acudir? No.',23,INK,True)
text(s,.65,5.95,5.91,.64,'Confirmar fuente, apertura, propósito y acceso antes de compartir el aviso.',16,GRAY)
rect(s,7.21,2.62,5.44,3.93,DARK,radius=True);tag(s,7.49,2.93,'SMS SIMULADO',True,2.32)
text(s,7.49,3.6,4.85,.83,'Ubicación + horario\n+ recepción / entrega de ayudas',20,WHITE,True)
picture(s,ASSETS/'sms.png',7.47,4.81,4.9,.87)
text(s,7.5,5.99,4.8,.34,'100 destinatarios ficticios · 0 envíos reales',12,MINT,True)
notes(s,'La ciudadanía aprende qué hacer y cómo distinguir los tipos de punto. El aviso se dirige a toda la población, sin identificar quién es damnificado. La captura muestra cien destinatarios ficticios, entregas y pendientes simulados. El borrador mantiene apertura no confirmada; no se invita a acudir a un lugar que solo fue seleccionado en el mapa.',20)

# 12 — closing, metrics framed as intended value not measured outcomes.
s=new('Valor público',dark=True)
text(s,.64,1.35,11.95,1.56,'Preparar el espacio.\nHacer visible lo que falta.',41,WHITE,True)
for j,(a,b) in enumerate([('CRITERIOS VISIBLES','Entender por qué se revisa un lugar.'),('OCUPACIÓN CONTROLADA','Detectar propuestas que exceden el área.'),('COMUNICACIÓN CLARA','Distinguir borrador de apertura confirmada.')]):
    x=.66+j*4.15;rule(s,x,3.75,3.71,'355965');text(s,x,4.1,3.73,.63,a,16,MINT,True);text(s,x,4.93,3.68,.82,b,18,WHITE)
text(s,.66,6.12,12,.35,BRAND['name']+' propone. La autoridad decide.',25,MINT,True)
small(s,.67,6.67,'William Ortiz · Herlin Echeverry · Pablo Arango · Daniel Celis',True,w=12)
notes(s,'La propuesta es hacer explicable la preparación: criterios visibles, ocupación controlada y comunicación clara. Son capacidades demostradas en un prototipo, no impacto medido en una emergencia. '+BRAND['name']+' propone; la autoridad decide.',15)

# 13 — evidence and technical architecture, concrete current scope.
s=new('Anexo · evidencia de producto','Qué está conectado hoy en el prototipo.')
for j,(a,b,c) in enumerate([
('DATOS PÚBLICOS','GeoJSON local: inventario, comunas y amenazas.','Fuente y fecha de corte conservadas.'),
('REGLAS EN EL NAVEGADOR','Filtros, área cartográfica y necesidades de referencia.','Criterios deterministas y resultados explicables.'),
('BORRADOR Y SEGUIMIENTO','Exportación JSON y estados locales de revisión.','No acredita recepción por una entidad.'),
('COMUNICACIÓN DE DEMO','Vista previa SMS y resultados agregados simulados.','Sin proveedor conectado ni teléfonos.')]):
    y=2.6+j*.92;tag(s,.64,y,a,w=3.22);text(s,4.09,y+.04,8.49,.36,b,17,INK,True);text(s,4.09,y+.49,8.49,.3,c,13,GRAY)
notes(s,'La arquitectura de esta etapa es un frontend con archivos públicos locales y reglas reproducibles. Los estados de revisión se guardan localmente y se exportan borradores. No hay integración operativa con proveedores SMS o entidades; las capturas de la presentación muestran la implementación real disponible.')

# 14 — operational safeguards, not a legal wall.
s=new('Anexo · controles de confianza','Lo desconocido sigue siendo desconocido.')
table(s,['CONTROL','CÓMO SE EXPRESA'],[
('Autoridad','Seleccionar o incluir en un borrador no activa un espacio.'),
('Seguridad técnica','No sustituye evaluación estructural ni conceptos de Bomberos.'),
('Datos mínimos','Sin nombres, documentos, teléfonos ni identificación individual.'),
('Explicabilidad','Criterios, fórmulas, fuentes y supuestos visibles.'),
('Simulaciones','Población, reservas, módulos y entregas SMS marcados SIMULADOS.'),
('Retorno','Revisión, checklist y acta como flujo previsto; no acta ejecutada.')
],[3.1,9.0],y=2.54,rowh=.53)
notes(s,'Estos controles limitan las afirmaciones del sistema. No presentamos el desconocimiento como cero, la cercanía como seguridad o una captura como autorización. El retorno al uso cotidiano es parte del proceso previsto, no un acto ya realizado. La validación con entidades sigue fuera de esta etapa.')

# 15 — roadmap with no invented dates or validated performance.
s=new('Anexo · siguiente etapa','Qué falta para una operación real.')
for j,(n,a,b) in enumerate([('01','MEDIR','Área útil, accesos, servicios y condiciones de atención.'),('02','COMPROBAR','Encaje de módulos, restricciones, carga y evaluación técnica.'),('03','INTEGRAR','Proveedor SMS, permisos, cobertura y trazabilidad operativa.'),('04','VALIDAR','Disponibilidad, responsables y coordinación con actores reales.')]):
    x=.65+j*3.12;circle(s,x,2.88,.52,TEAL);text(s,x+.06,3,.4,.27,n,13,WHITE,True,PP_ALIGN.CENTER)
    if j<3:arrow(s,x+.7,3.15,x+2.75,3.15)
    text(s,x,3.81,2.83,.4,a,18,TEAL,True);text(s,x,4.46,2.73,1.44,b,18,INK)
rect(s,.65,6.27,12,.48,PALE,radius=True);text(s,.83,6.39,11.65,.25,'Pendientes de operación, no requisitos realizados. La validación con entidades está fuera de la etapa actual.',12,INK)
notes(s,'Esta hoja de ruta no contiene fechas o resultados inventados. Para operar hacen falta mediciones, comprobaciones técnicas, integración y validación. En la etapa actual el usuario excluyó la validación con entidades; se conserva como pendiente, no como logro. Los objetivos TRL e IRL del hackathon no se presentan como certificaciones alcanzadas.')

# 16 — primary sources and actual implementation trail.
s=new('Anexo · fuentes','La propuesta se puede rastrear hasta sus datos.')
rows=[
('Inventario y comunas','Alcaldía de Cali · DAPM / IDESC · archivos públicos del proyecto','https://datos.cali.gov.co/dataset/epou-espacio-publico-efectivo'),
('Cartografía de comunas','Comunas de Santiago de Cali · atribución CC BY','https://www.datos.gov.co/dataset/Comunas-de-Santiago-de-Cali/dx2g-2mhm'),
('Necesidades de referencia','Manual Esfera 2018 · valores de planificación a contextualizar','https://spherestandards.org/wp-content/uploads/Sphere-Handbook-2018-EN.pdf'),
('Orientación ciudadana','Recomendaciones UNGRD consultadas en el proyecto','https://portal.gestiondelriesgo.gov.co/Paginas/inundaciones.aspx'),
('Código y bitácora','docs/PROYECTO.md · funciones y pruebas del frontend','https://github.com/leonidas452528/hackaton-camaracomercio2026')]
for j,(title,desc,url) in enumerate(rows):
    y=2.45+j*.76;text(s,.65,y,3.1,.34,title,14,TEAL,True);text(s,3.72,y,8.9,.31,desc,13,INK)
    sh=text(s,3.72,y+.38,8.9,.22,url,8.8,TEAL);sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url
small(s,.65,6.57,'Corte 25 sep 2026 · Adaptaciones cartográficas CC BY-SA 4.0 · Ilustración conceptual IA identificada · Sin datos personales',w=12)
notes(s,'Fuentes primarias y trazabilidad. Las imágenes de cartografía derivan del inventario local del proyecto. Las capturas son del frontend con la marca Territorio Preparado. La imagen arquitectónica es una ilustración conceptual generada para esta presentación: no una fotografía de Cali ni un diseño aprobado.')

# Remove inherited template shadows so diagrams remain crisp in PowerPoint and PDF.
from pptx.oxml.ns import qn
for slide in prs.slides:
    for shape in slide.shapes:
        for effect in shape._element.findall('.//'+qn('a:effectRef')):effect.set('idx','0')
prs.core_properties.title=BRAND['name']+' — '+BRAND['tagline']
prs.core_properties.author='Equipo '+BRAND['name']
prs.core_properties.subject='Presentación ampliada del sistema de preparación de espacios'
prs.save(OUT/'Presentacion_Ampliada.pptx')
(OUT/'src/guion.json').write_text(json.dumps(SLIDES,ensure_ascii=False,indent=2)+'\n')
(OUT/'Guion_del_pitch.md').write_text('# '+BRAND['name']+' · Guion ampliado\n\n12 diapositivas principales: 240 segundos sugeridos. Cuatro anexos para preguntas.\n\n'+'\n\n'.join(f'## Diapositiva {r["slide"]}'+(f' · {r["seconds"]} segundos' if r['seconds'] else ' · Anexo')+'\n\n'+r['notes'] for r in SLIDES)+'\n')
print('Presentación ampliada:',len(prs.slides),'diapositivas;',sum(r['seconds'] or 0 for r in SLIDES),'segundos principales.')
