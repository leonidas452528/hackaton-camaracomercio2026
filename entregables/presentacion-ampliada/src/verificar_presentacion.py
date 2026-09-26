from pathlib import Path
import json,re
import pymupdf as fitz
from pptx import Presentation
from PIL import Image,ImageOps,ImageDraw
ROOT=Path(__file__).resolve().parents[1]
doc=fitz.open(ROOT/'Presentacion_Ampliada.pdf');prs=Presentation(ROOT/'Presentacion_Ampliada.pptx')
assert len(doc)==len(prs.slides)==16
problems=[]
norm=lambda s: re.sub(r'\s+','',s)
for i,(page,slide) in enumerate(zip(doc,prs.slides)):
    actual=norm(page.get_text())
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text.strip():
            if norm(shape.text) not in actual:problems.append({'slide':i+1,'text':shape.text})
        if shape.left<0 or shape.top<0 or shape.left+shape.width>prs.slide_width+500 or shape.top+shape.height>prs.slide_height+500:
            problems.append({'slide':i+1,'outsideCanvas':shape.name})
    assert abs(page.rect.width/page.rect.height-16/9)<.001
    assert slide.has_notes_slide
    pix=page.get_pixmap(matrix=fitz.Matrix(1.33333,1.33333),alpha=False)
    pix.save(str(ROOT/f'assets/slide-{i+1:02}.png'))
for half in range(3):
    sheet=Image.new('RGB',(1280,1080),'#D5DEDA')
    for k in range(6):
        if half*6+k>=len(doc):break
        im=Image.open(ROOT/f'assets/slide-{half*6+k+1:02}.png').resize((632,355))
        sheet.paste(im,((k%2)*640,(k//2)*360))
    sheet.save(ROOT/f'Vista_general_{half+1}.jpg',quality=94)
Image.open(ROOT/'assets/slide-01.png').save(ROOT/'Portada.jpg',quality=95)
result={'slides':16,'mainSlides':12,'appendixSlides':4,'aspectRatio':'16:9','speakerNotes':sum(s.has_notes_slide for s in prs.slides),'mainPitchSeconds':sum(r['seconds'] or 0 for r in json.loads((ROOT/'src/guion.json').read_text())),'editableTextBoxes':sum(1 for s in prs.slides for sh in s.shapes if sh.has_text_frame and sh.text.strip()),'textOrBoundsProblems':problems,'pdfBytes':(ROOT/'Presentacion_Ampliada.pdf').stat().st_size,'pptxBytes':(ROOT/'Presentacion_Ampliada.pptx').stat().st_size}
(ROOT/'verificacion.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(result,ensure_ascii=False,indent=2))

assert not problems, "Hay problemas de texto o límites del lienzo"
