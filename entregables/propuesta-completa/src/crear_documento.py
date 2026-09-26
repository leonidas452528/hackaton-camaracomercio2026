"""Documento editable de la propuesta, con alcance contrastado con la bitácora."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'entregables/propuesta-completa'
ASSETS=ROOT/'entregables/presentacion-ampliada/assets'
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.72)
sec.left_margin=sec.right_margin=Inches(.78)
sec.header_distance=sec.footer_distance=Inches(.3)
styles=doc.styles
for name in ['Normal','Body Text','List Bullet','List Number']:
 st=styles[name];st.font.name='DejaVu Sans';st.font.size=Pt(10.5)
 st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.1
for name,size in [('Title',34),('Heading 1',24),('Heading 2',14),('Heading 3',11)]:
 st=styles[name];st.font.name='DejaVu Sans';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('14323D');st.font.bold=True
 st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(9)
styles['Caption'].font.size=Pt(8);styles['Caption'].font.color.rgb=RGBColor.from_string('536D70')
header=sec.header.paragraphs[0];header.text='TP  /  TERRITORIO PREPARADO';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Propuesta · 25 septiembre 2026   |   ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
sec.different_first_page_header_footer=True
PAGES=[]
def p(s,style=None):return doc.add_paragraph(s,style)
def h(s):doc.add_heading(s,2)
def bullets(items):
 for item in items:p(item,'List Bullet')
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
 if widths:
  for col,w in zip(t.columns,widths):col.width=Inches(w)
 for cell,text in zip(t.rows[0].cells,headers):
  cell.text=text
  tcPr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'14323D');tcPr.append(sh)
  for run in cell.paragraphs[0].runs:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for cell,text in zip(cells,row):
   cell.text=text
   if i%2==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EFF5F4');cell._tc.get_or_add_tcPr().append(sh)
 for row in t.rows:
  flag=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(flag)
  for cell in row.cells:
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(5);par.paragraph_format.space_before=Pt(5)
    for run in par.runs:run.font.size=Pt(9)
 p('')
 return t

def page(title,kicker):
 if PAGES:doc.add_page_break()
 PAGES.append(title);p(kicker.upper(),'Subtitle');doc.add_heading(title,1)
def image(name,width,caption):
 doc.add_picture(str(name),width=Inches(width));p(caption,'Caption')

page('Territorio Preparado','Propuesta integral · Documento para presentar')
p('Espacios y ayudas coordinados ante emergencias','Subtitle')
p('Planeación del uso temporal de espacios públicos de Cali, con criterios visibles, balance de superficies y necesidades por resolver.')
image(ASSETS/'espacio-conceptual.png',6.55,'Ilustración conceptual generada con IA para la presentación. No representa un parque real, obras existentes ni un diseño técnico aprobado.')
h('La propuesta en una frase')
p('Ayudar a preparar dónde y cómo prestar apoyo a la población ante una emergencia, convirtiendo datos territoriales y supuestos explícitos en una propuesta revisable. La autoridad decide la activación.')
p('Hackathon Smart City Expo Cali 2026\nRETO-01 Cali Activa · Cámara de Comercio de Cali y Alcaldía de Santiago de Cali')
p('William Ortiz · Herlin Echeverry\nPablo Arango · Bryan Martínez Villamarín')
p('Corte de esta entrega: 25 de septiembre de 2026.\nPrototipo demostrable; no acredita operación ni validación con entidades.','Caption')

page('1. Resumen ejecutivo','Propósito y problema')
p('Un parque identificado en un mapa no basta para organizar una respuesta. Antes de proponerlo como punto de acopio o atención hay que revisar su exposición, el espacio que puede ocuparse, los accesos y los servicios necesarios. La información incompleta debe quedar visible.')
h('Objetivo general')
p('Apoyar la preparación del uso temporal de espacios públicos de Cali ante distintas amenazas, con una propuesta explicable que conecte ubicación, restricciones, dimensionamiento, necesidades y comunicación.')
h('El resultado central')
p('Un borrador por espacio: identificación y fuente, escenario, criterios de comparación, reservas de área, necesidades calculadas, datos faltantes y seguimiento local. Este borrador permite revisar la propuesta antes de una decisión externa de activación.')
table(['Problema de preparación','Solución planteada'],[
('Ubicación sin caracterización suficiente','Consultar inventarios públicos y mostrar evidencia y vacíos.'),
('Un mismo lugar se considera útil para toda emergencia','Cambiar la revisión según la amenaza y sus restricciones.'),
('Se ocupa el parque sin reservar accesos o atención','Descontar superficies y detectar exceso de almacenamiento propuesto.'),
('No queda claro qué hace falta ni quién debe revisarlo','Estimar necesidades y registrar responsables propuestos y pendientes.'),
('La población puede confundir un punto en el mapa con un lugar abierto','Educar y comunicar condiciones; exigir apertura confirmada antes de invitar a acudir.')],[2.8,3.8])
h('Alcance de esta entrega')
p('Se presenta una aplicación local, una demostración de escenarios, una maqueta conceptual y material de pitch. El sistema no modela físicamente la propagación de una catástrofe. Demuestra la lógica de preparación de la atención a la población.')
p('Lectura: fases y actores (sección 2); preparación a retorno (3–7); tecnología (8); demo (9); evidencia (10); hoja de ruta (11); fuentes y entregables (12).','Caption')

page('2. Solución y fases completas','Del territorio a la recuperación')
p('La propuesta reúne cinco fases del ciclo de atención. Su implementación es parcial y se declara en cada fase para no confundir la visión completa con una operación ya habilitada.')
table(['Fase','Resultado esperado','Estado actual'],[
('1. Conocer y prevenir','Ficha del espacio, restricciones y vacíos.','Consulta territorial y educación implementadas.'),
('2. Dimensionar y preparar','Propuesta de ocupación y necesidades.','Cálculos implementados con supuestos SIMULADOS.'),
('3. Coordinar y decidir','Pendientes asignados y decisión de activación.','Seguimiento local; decisión e integración institucional pendientes.'),
('4. Informar y acompañar','Aviso claro y seguimiento de la operación.','Aviso y envío SMS SIMULADOS; operación real pendiente.'),
('5. Recuperar y aprender','Inspección, devolución y actualización de la ficha.','Representación conceptual; checklist y acta operativos pendientes.')],[1.55,2.4,2.65])
h('Actores y responsabilidades previstas')
bullets([
'Equipo de preparación: consulta el territorio, registra supuestos, revisa alternativas y entrega el borrador.',
'Autoridad competente: determina la activación y sus condiciones. Secretaría de Gestión del Riesgo o Consejo Municipal de Gestión del Riesgo.',
'Entidades de servicios: revisión y atención de necesidades según sus competencias; las asignaciones del prototipo son propuestas, no compromisos confirmados.',
'Comunidad: aprende a distinguir funciones de los puntos y consulta información de apertura y acceso. La difusión prevista es general; no exige identificar damnificados.'
])
h('Cadena de decisión')
p('Escenario → evidencia → candidato → superficie → necesidades → borrador → revisión de autoridad → comunicación de apertura → operación → retorno.')
p('Si falta evidencia se conserva el pendiente; si hay restricción se revisa otro espacio; si excede el área se ajusta la propuesta. La presencia en el mapa nunca equivale a apertura.')

page('3. Conocer el territorio y prevenir','Fase 1 · Implementada con límites de evidencia')
p('El usuario consulta espacios reales por comuna o barrio, abre su ficha y revisa ubicación, procedencia y cruces disponibles. El inventario de espacio público efectivo contiene 1.970 registros [1]. Estos datos describen el territorio, no una lista de alojamientos habilitados.')
h('Soluciones disponibles')
bullets([
'Mapa con polígonos y puntos diferenciados, filtros y fichas consultables. Los puntos sin contorno no reciben un área inventada.',
'Enlaces por coordenadas a Google Maps, satélite y Street View. Las calles dentro de la aplicación son de OpenStreetMap [3].',
'Comparación de hasta tres polígonos del sector por restricciones conocidas y distancia en línea recta. No es distancia por ruta ni ranking de capacidad operativa.',
'Orientación preventiva antes, durante y después para inundación, sismo e incendio forestal. Incluye aprendizaje sobre apertura y tipos de puntos [5].'
])
table(['Amenaza','Regla actual','Límite'],[
('Inundación','Descarta cruces con capas disponibles.','Sin cruce detectado no significa ausencia de riesgo.'),
('Sismo','Excluye cruces de licuación/corrimiento; exige inspección aun sin cruce.','No constituye una evaluación sísmica integral.'),
('Incendio forestal','No preselecciona sin evidencia del incidente y restricciones.','No predice fuego, humo ni perímetros seguros.'),
('Incendio en edificación','Muestra preparación y condiciones pendientes.','No emite conceptos de seguridad humana o Bomberos.'),
('Sequía','No produce candidatos sin evidencia suficiente.','Faltan datos de amenaza y abastecimiento.')],[1.3,2.7,2.6])
p('La población del escenario modifica las necesidades calculadas; no modifica el orden de cercanía, porque no hay capacidades confirmadas para realizar ese ajuste.')
p('Salida: espacio identificado, criterios visibles y lista de evidencia pendiente. Fuentes [1], [2] y catálogo de datos del proyecto.','Caption')

page('4. Dimensionar la propuesta','Fase 2 · Superficie, población y necesidades')
p('Una vez elegido el registro, el sistema dibuja su huella pública y permite completar un balance de acopio. Las reservas introducidas son supuestos SIMULADOS / por verificar, no mediciones automáticas del parque.')
h('Regla de superficie')
p('Área para almacenamiento = área cartográfica − área no utilizable − circulación y accesos − atención y espera − otros usos.')
p('Máximo teórico de módulos = parte entera del área restante / huella de cada módulo. Las reservas no deben superponerse entre sí. Sin datos completos o sin polígono no se calcula una capacidad utilizable.')
h('Controles incorporados')
bullets([
'Rechazo de valores inválidos y campos vacíos; un dato faltante no equivale a cero.',
'Reservas explícitas para circulación y atención; aviso si la propuesta supera la superficie.',
'Inclusión de la distribución en el borrador solo con balance válido y condiciones de amenaza que no excluyan ni dejen pendiente al espacio.',
'Cambiar medidas, espacio o amenaza invalida la inclusión anterior y obliga a revisarla.'
])
h('Necesidades de atención')
table(['Referencia [4]','Cálculo','Escenario SIMULADO: 300 personas'],[
('Agua para necesidades básicas','15 L por persona y día','4.500 L/día'),
('Baños','Redondear hacia arriba personas / 20','15 unidades'),
('Superficie habitable cubierta','3,5 m² por persona','1.050 m²')],[2.05,2.15,2.4])
p('Son referencias humanitarias que deben contextualizarse. El área cartográfica del parque no equivale a superficie cubierta. Las necesidades calculadas no son existencias ni garantizan suministro.')
p('Si no se conocen baños, agua o área cubierta disponibles, la brecha queda no calculable. El balance de acopio no determina aforo de personas, toneladas almacenables ni encaje físico en polígonos irregulares.')

page('5. Coordinar y someter a revisión','Fase 3 · Seguimiento local y decisión externa')
p('La preparación reúne necesidades, supuestos y faltantes en un mismo borrador. La utilidad es que el equipo pueda explicar qué propone y qué debe resolverse; una selección en pantalla no activa un espacio.')
table(['Servicio','Referencia en el prototipo','Estado de coordinación'],[
('Agua básica','EMCALI, responsable propuesto','Disponibilidad y aceptación no verificadas.'),
('Baños y saneamiento','UAESP, responsable propuesto','Requerimientos y operación por comprobar.'),
('Área cubierta y preparación','Gestión del Riesgo, responsable propuesto','Condiciones y coordinación por revisar.')],[1.5,2.4,2.7])
h('Lo que puede hacer el usuario hoy')
bullets([
'Consultar necesidades de referencia y distinguir existencias desconocidas.',
'Cambiar estados locales de revisión por espacio y servicio: “Por medir” y “En revisión”.',
'Descargar un JSON con escenario, candidatos, criterios, necesidades, reservas, balance de acopio, seguimiento y fuentes.'
])
h('Proceso previsto antes de una apertura')
p('La autoridad revisaría el borrador junto con las evaluaciones y verificaciones correspondientes. Deben resolverse condiciones de uso, acceso, servicios y coordinación. La decisión y su evidencia se incorporarán en una etapa operativa posterior; hoy no hay aprobación electrónica ni recepción confirmada por entidades.')
h('Condiciones de confianza')
p('El sistema recomienda y conserva la explicación de sus reglas. No sustituye evaluaciones estructurales ni conceptos técnicos de seguridad humana o contra incendios. No emplea nombres, documentos, teléfonos ni reconocimiento facial. La información del escenario es agregada y simulada.')
p('La validación con entidades está fuera de la etapa actual por decisión del usuario. No se presenta como realizada ni como requisito para ejecutar esta demo. Un cambio de estado local tampoco acredita cierre de una brecha.')

page('6. Informar y acompañar','Fase 4 · Cultura ciudadana y difusión')
p('La comunicación prevista se dirige a la población general: personas que necesitan ayuda, quienes pueden aportar y quienes necesitan conocer las condiciones de atención. No se presupone saber quién fue damnificado.')
h('Funciones de cada punto')
table(['Tipo','Propósito','Información que debe verificarse'],[
('Acopio','Recibir, almacenar y distribuir ayudas.','Qué recibe o entrega, horario, capacidad y apertura.'),
('Alojamiento temporal','Prestar alojamiento y servicios de atención.','Condiciones de habitabilidad, servicios y autorización.'),
('Encuentro','Reunión u orientación según el plan correspondiente.','Ubicación, acceso e instrucciones oficiales.')],[1.4,2.2,3.0])
h('Canal SMS de demostración')
p('Recorrido: mapa → seleccionar un espacio → Preparar este espacio → Preparar aviso SMS de acopio → Simular envío masivo.')
p('El borrador incluye nombre/ID real, ubicación enlazada a Google Maps y campos propuestos de horario y recepción/entrega. Se puede revisar, copiar y descargar. La apertura permanece no confirmada y el mensaje no invita a acudir a un lugar sin habilitación.')
p('La simulación permite audiencia agregada ficticia, procesamiento y resultados completos o con pendientes. Ejemplo SIMULADO: 100 destinatarios, 95 entregas y 5 pendientes. Envíos reales: cero. No se usan teléfonos, proveedor o credenciales; no se afirma cobertura de toda Cali.')
h('Orientación preventiva')
p('La sección educativa enseña a revisar fuente, fecha, función del punto, apertura y acceso. Incluye una lista de aprendizaje y una pregunta interactiva que diferencia inventario de habilitación. No es una alerta vigente ni una orden de evacuación [5].')
h('Acompañamiento de operación: siguiente etapa')
p('Se propone incorporar información actualizada de ocupación y servicios, gestión de incidencias y avisos autorizados de cambios. Estas funciones requieren datos y coordinación operativa; la demo actual no monitorea ocupación real ni entregas reales de ayudas.')

page('7. Recuperar el espacio y aprender','Fase 5 · Proceso propuesto')
p('La transformación temporal debe terminar con el regreso del espacio a su uso cotidiano. La propuesta incluye recuperación para que la preparación no se limite a abrir un punto de atención.')
table(['Paso previsto','Evidencia que se necesitaría'],[
('Cerrar o trasladar la atención','Decisión de autoridad y comunicación de cierre o traslado.'),
('Retirar población, ayudas y equipos','Registro agregado de cierre y destino verificado del inventario.'),
('Limpiar y gestionar residuos','Registro de actividades y revisión de condiciones.'),
('Inspeccionar y reparar','Evaluación competente y pendientes identificados.'),
('Entregar el espacio','Checklist revisado y acta con responsables.'),
('Actualizar la caracterización','Lecciones aprendidas y datos de preparación corregidos.')],[2.5,4.1])
h('Representación actual')
p('La maqueta permite visualizar uso cotidiano, emergencia y recuperación. Es una explicación del ciclo, no un registro de una operación realizada. No existe una inspección final ni un acta operativa verificadas en esta entrega.')
h('Cómo debería cerrar el ciclo')
p('La salida futura sería una ficha actualizada del espacio con estado de retorno, pendientes y evidencia de entrega. Debería conservar el historial de cambios y permitir que una preparación posterior use información revisada.')
h('Criterio de avance')
p('La representación visual no cierra administrativamente el lugar. El retorno debe depender de revisión competente y evidencias. El checklist y acta son componentes por implementar y validar cuando se retome la etapa operativa.')
p('Esta fase completa la visión antes–durante–después del reto, sin atribuir al prototipo una gestión de recuperación que aún no ejecuta.')

page('8. Maqueta y arquitectura del prototipo','Cómo se explica y cómo funciona')
p('La aplicación utiliza Vite, React y TypeScript; Leaflet para el mapa; Three.js / React Three Fiber para la visualización. Las reglas actuales son deterministas y explicables. No se presenta un modelo predictivo de inteligencia artificial.')
table(['Componente','Implementación y límite'],[
('Datos territoriales','Archivos públicos locales con procedencia. No es una transmisión de incidentes en tiempo real.'),
('Reglas de preparación','Cruces disponibles, distancias, balance de superficie y necesidades.'),
('Persistencia y exportación','Seguimiento en el navegador y borrador JSON. Sin coordinación compartida de usuarios.'),
('Servicios externos','Calles OpenStreetMap y enlaces Google Maps; SMS sin proveedor real.')],[1.7,4.9])
h('Qué se adapta y qué se ilustra')
p('La selección de un polígono adapta su huella y el encuadre. No coloca automáticamente refugios, cisternas o almacenamiento. Los puntos sin contorno quedan sin planta calculable. La escena general muestra una solución conceptual de refugio, acopio, circulación y agua.')
image(ROOT/'maqueta3d/deliverables/renders/sistema_agua.png',5.55,'Captura de la maqueta del proyecto: circuito de agua ilustrativo, no instalación verificada ni simulación hidráulica.')
p('El concepto de agua muestra captación, conducción, almacenamiento y servicio no potable. Capacidad útil, caudal, energía, tratamiento, suministro y descarga siguen por verificar. No se conecta esta ilustración con una garantía de los 15 L/persona/día calculados para necesidades básicas.')

page('9. Demostración para presentar','Caso reproducible · Datos reales y supuestos visibles')
p('Caso territorial: epou-9465, Parque · Colinas del Sur · EPE_1064. Huella cartográfica real: 447,2149007228903 m² [1]. Las reservas y módulos siguientes son SIMULADOS; no constituyen medición del área útil.')
table(['Entrada del balance','Valor de demostración'],[
('Área no utilizable','100 m² · SIMULADO'),
('Circulación y accesos','80 m² · SIMULADO'),
('Atención y espera','60 m² · SIMULADO'),
('Otros usos / ocupación','40 m² · SIMULADO'),
('Área restante','167,21 m² · resultado del supuesto'),
('Huella por módulo','10 m² · SIMULADO'),
('Proponer 20 módulos','200 m²: excede 32,79 m²'),
('Ajustar a 16 módulos','160 m²: restan 7,21 m²')],[3.3,3.3])
h('Guion de recorrido')
for s in [
'Abrir el mapa y buscar epou-9465. Mostrar su ficha, fuente y enlace a Google Maps.',
'Pulsar Preparar este espacio. Explicar la huella y elegir un escenario; leer las restricciones antes de proponer uso.',
'Introducir 300 personas SIMULADAS y mostrar necesidades, separándolas de las existencias desconocidas.',
'Completar las reservas y comparar 20 módulos con 16. Si la amenaza excluye o deja pendiente el lugar, explicar que un balance favorable no habilita su inclusión.',
'Descargar el borrador. Mostrar prevención y Simular envío masivo, remarcando cero envíos reales y apertura no confirmada.'
]:p(s,'List Number')
p('Resultado que se puede demostrar: el software detecta exceso por área, conserva vacíos y permite revisar la propuesta. No demuestra que quepan físicamente 16 módulos, ni un aforo de 16 familias, ni un parque abierto.')

page('10. Evidencia y alcance de la entrega','Lo que tenemos para presentar')
table(['Disponible','Evidencia o salida'],[
('Explorador territorial','Inventarios, filtros, fichas, cruces y enlaces por coordenadas.'),
('Preparación por espacio','Escenario, huella, candidatos, necesidades y borrador JSON.'),
('Control de acopio','Reservas explícitas, exceso, ajuste e invalidación al editar.'),
('Visualización 3D','Refugio, acopio, agua y ciclo de uso ilustrativos; exportación GLB.'),
('Prevención y difusión','Aprendizaje y mensajes con envío SMS SIMULADO.'),
('Presentación del pitch','PowerPoint editable y PDF: 12 láminas principales + 4 anexos; notas para cuatro minutos sugeridos.'),
('Documentación','Bitácora, fuentes, alcance real, guía GitHub y este documento Word.')],[2.2,4.4])
h('Verificación registrada')
p('La revisión de la entrega anterior registró compilación de producción, 22 pruebas unitarias y tres pruebas de navegador de avisos, SMS y prevención móvil. La presentación se verificó con 16 páginas, notas y 288 cuadros de texto editables, sin texto perdido ni elementos fuera del lienzo [6]. Estas pruebas son evidencia del software, no de disponibilidad física ni validación institucional.')
h('Qué no afirmamos como logrado')
bullets([
'Operación real, espacios habilitados, aforos certificados o servicios disponibles.',
'Distribución automática y encaje geométrico de kits; diseño estructural o hidráulico aprobado.',
'Gestión real de lotes y trazabilidad de ayudas: la representación no contiene inventario verificado.',
'Envíos SMS reales, cobertura de la población, alertas automáticas o rutas seguras.',
'Coordinación validada con entidades, impacto medido o nivel de madurez certificado.'
])
p('TRL 3 e IRL 3 son metas del proyecto documentadas en la bitácora. No se presentan como acreditaciones obtenidas; la validación con actores permanece pendiente y fuera de esta etapa.')

page('11. Ruta de desarrollo y sostenimiento','Qué sigue y cómo comprobarlo')
table(['Etapa propuesta','Trabajo y condición de salida'],[
('Consolidar la demo','Mantener fuentes, supuestos y flujo de preparación coherentes; poder repetir el caso de superficie y exportar su borrador.'),
('Medir','Levantar área útil, accesos, obstáculos y servicios. Salida: fichas con evidencia y fecha de revisión.'),
('Comprobar distribución','Verificar encaje de módulos y condiciones físicas con personal competente. Salida: propuesta revisada, sin inferir aforo por área bruta.'),
('Integrar operación','Backend, permisos, historial de evidencias y proveedor de difusión. Salida: pruebas de persistencia, autorización y entrega.'),
('Validar con actores','Revisar responsabilidades, disponibilidad y coordinación. Se retoma cuando el equipo abra esa etapa; hoy está excluida.'),
('Pilotar y cerrar el ciclo','Ejercicio autorizado con indicadores, revisión de incidencias y retorno documentado.')],[1.9,4.7])
h('Indicadores propuestos, todavía sin medición')
bullets([
'Tiempo para producir un borrador explicable en un escenario controlado.',
'Porcentaje de fichas con evidencia vigente de área útil y servicios.',
'Casos de sobreocupación detectados frente a casos de prueba definidos.',
'Pendientes con responsable confirmado y evidencia de atención.',
'Comprensión ciudadana de apertura y función del punto; entrega real de avisos cuando exista integración.'
])
h('Recursos y continuidad')
p('Se requiere trabajo geográfico y de campo, desarrollo web, revisión técnica, coordinación y comunicación. Los costos de hosting, SMS, medición, equipamiento y mantenimiento deben cotizarse; no existe en esta entrega un presupuesto aprobado ni financiación confirmada.')
p('El mantenimiento propuesto incluye fechas de actualización de datasets, revisión de licencias, respaldos, trazabilidad de cambios y responsabilidades de actualización. No se promete operación sin conexión completa: datos locales, calles externas y enlaces tienen dependencias distintas.')

page('12. Fuentes y archivos de entrega','Trazabilidad para el equipo y los evaluadores')
p('Documento elaborado con la bitácora y el alcance implementado del repositorio. No incorpora nuevos datasets, cifras de emergencia ni afirmaciones de impacto. Las fuentes se conservan con su procedencia y licencias.')
refs=[
('[1] Espacio público efectivo — DAPM / IDESC','https://datos.cali.gov.co/dataset/epou-espacio-publico-efectivo'),
('[2] Comunas de Santiago de Cali — DAPM','https://www.datos.gov.co/dataset/Comunas-de-Santiago-de-Cali/dx2g-2mhm'),
('[3] Calles y atribución — OpenStreetMap','https://www.openstreetmap.org/copyright'),
('[4] Referencias humanitarias — Manual Esfera 2018','https://spherestandards.org/wp-content/uploads/Sphere-Handbook-2018-EN.pdf'),
('[5] Orientación educativa — UNGRD / fuentes archivadas en bitácora','https://portal.gestiondelriesgo.gov.co/Paginas/recomendaciones-incendios-forestales.aspx'),
('[6] Evidencia local y código del proyecto','docs/PROYECTO.md; docs/PROPUESTA_REAL.md; maqueta3d/src/planning.ts; maqueta3d/src/calculateAcopioCapacity.ts; entregables/presentacion-ampliada/verificacion.json')]
for title,url in refs:
 h(title);p(url,'Caption')
p('Atribución cartográfica: Alcaldía de Santiago de Cali – DAPM / IDESC. Adaptaciones del inventario: CC BY-SA 4.0; comunas de origen: CC BY. Catálogo completo de descargas y licencias: prototipo/datos/FUENTES.md y fuentes.csv. Corte local: 25 de septiembre de 2026.','Caption')
h('Paquete para presentar')
bullets([
'Propuesta completa: entregables/propuesta-completa/Territorio_Preparado_Propuesta_Completa.docx y PDF complementario.',
'Pitch: entregables/presentacion-ampliada/Presentacion_Ampliada.pptx y PDF; Guion_del_pitch.md.',
'Prototipo: carpeta maqueta3d. Ejecutar npm ci y npm run dev dentro de esa carpeta; abrir http://localhost:5173.',
'Alcance resumido: docs/PROPUESTA_REAL.md. Bitácora: docs/PROYECTO.md y PDF. Publicación: docs/SUBIR_A_GITHUB.md.'
])
p('La ilustración de portada fue generada con image_gen para la presentación ampliada; prompt y procedencia en su carpeta src/prompt-ilustracion.json. La captura del agua procede de la maqueta del proyecto. Ambas tienen alcance conceptual, sin inventario ni obras certificadas.','Caption')
p('Mensaje de cierre: preparar el espacio, hacer visible lo que falta y entregar una propuesta que la autoridad pueda revisar.')

# Texto alternativo de las imágenes y metadatos del documento.
for i,shape in enumerate(doc.inline_shapes):
 shape._inline.docPr.set('descr', ['Ilustración conceptual de acopio, atención, accesos y servicios; no parque real.','Captura de la maqueta: circuito conceptual de agua no potable.'][i])
doc.core_properties.title='Territorio Preparado — Propuesta completa'
doc.core_properties.subject='Fases, soluciones, prototipo demostrable y alcance pendiente'
doc.core_properties.author='William Ortiz, Herlin Echeverry, Pablo Arango y Bryan Martínez Villamarín'
doc.core_properties.keywords='Cali, espacios públicos, preparación, acopio, emergencias, prototipo'
doc.save(OUT/'Territorio_Preparado_Propuesta_Completa.docx')
(OUT/'src/secciones.json').write_text(json.dumps(PAGES,ensure_ascii=False,indent=2)+'\n')
print(f'Documento creado: {len(PAGES)} secciones paginadas.')
