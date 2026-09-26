from pathlib import Path
import json,re
import pymupdf
from docx import Document
from PIL import Image,ImageOps,ImageDraw
OUT=Path(__file__).resolve().parents[1]
doc=Document(OUT/'Territorio_Preparado_Propuesta_Completa.docx')
pdf=pymupdf.open(OUT/'Territorio_Preparado_Propuesta_Completa.pdf')
norm=lambda s:re.sub(r'\s+','',s).replace('\u00ad','')
actual=norm(''.join(p.get_text() for p in pdf))
texts=[p.text for p in doc.paragraphs if p.text.strip()]
texts += [p.text for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs if p.text.strip()]
missing=[s for s in texts if norm(s) not in actual]
pageinfo=[]
for i,page in enumerate(pdf):
 text=page.get_text();pageinfo.append({'page':i+1,'characters':len(text),'firstLines':text.splitlines()[:4]})
 page.get_pixmap(matrix=pymupdf.Matrix(.9,.9),alpha=False).save(OUT/f'pagina-{i+1:02}.png')
for group in range((len(pdf)+5)//6):
 sheet=Image.new('RGB',(1050,1000),'#DDE6E5')
 for j in range(6):
  k=group*6+j
  if k>=len(pdf):break
  im=Image.open(OUT/f'pagina-{k+1:02}.png');im.thumbnail((338,478));x=(j%3)*350+(350-im.width)//2;y=(j//3)*500
  sheet.paste(im,(x,y));ImageDraw.Draw(sheet).text((x,y+479),f'Página {k+1}',fill='#14323D')
 sheet.save(OUT/f'Vista_general_{group+1}.jpg',quality=90)
report={'pages':len(pdf),'expectedPages':13,'tables':len(doc.tables),'images':len(doc.inline_shapes),'missingText':missing,'pageInfo':pageinfo}
(OUT/'verificacion.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False,indent=2))
assert not missing, 'Texto ausente en PDF'
assert len(pdf)==13, 'Revisar desbordes de página'
